import sys, os, re
import numpy as np
import pandas as pd
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
for p in [str(BASE) + '/libraries/CausalBootstrapping', str(BASE) + '/libraries/GRAPL',
          str(BASE) + '/libraries/causal_assistant', str(BASE) + '/tCB', str(BASE) + '/tests']:
    if p not in sys.path:
        sys.path.insert(0, p)

import causal_assistant as ca
from scipy.stats import gaussian_kde
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.preprocessing import StandardScaler as _SS

import transport_u_age as age

RANDOM_SEED = 42
U_DIMS = 2
N_AGE_BINS = 6
CB_FIT_METHOD = "kde"
META_COLS = {"Unnamed: 0", "tot_name"}
NV_FEAT = str(BASE) + '/data/neurovoz/extracted_features'
INTENSITY_IDX = slice(0, 38)

AGE_GRAPH = "y;X;Age; y->X; Age->X; Age->y;"
UPDRS_EDGES = [0, 10, 20, 40, 70, 200]
GRBAS_EDGES = [0, 1, 2, 3, 4]
GRAPH_MODELS = {
    "age":   "y;X;Age; y->X; Age->X; Age->y;",
    "grbas": "y;X;Age;GRBAS; y->X; Age->X; Age->y; y->GRBAS; GRBAS->X;",
    "updrs": "y;X;Age;UPDRS; y->X; Age->X; Age->y; y->UPDRS; UPDRS->X;",
    "both":  "y;X;Age;UPDRS;GRBAS; y->X; Age->X; Age->y; y->UPDRS; UPDRS->X; "
             "y->GRBAS; GRBAS->X;",
    "g-as[a]":   "y;X;Age;Sex;GRBAS; y->X; Age->X; Age->y; Sex->X; Sex->y; "
                 "y->GRBAS; GRBAS->X;",
    "ug-as[a]":  "y;X;Age;Sex;UPDRS;GRBAS; y->X; Age->X; Age->y; Sex->X; Sex->y; "
                 "y->UPDRS; UPDRS->X; y->GRBAS; GRBAS->X;",
    "xy":    "y;X; X->y;",
}
INTERVALS = {"-2": (-10, -6), "-1": (-5, -1), "1": (0, 5), "2": (6, 10)}

def _balanced_acc(y, pred):
    p0 = np.mean(pred[y == 0] == 0) if np.any(y == 0) else np.nan
    p1 = np.mean(pred[y == 1] == 1) if np.any(y == 1) else np.nan
    return 0.5 * (p0 + p1)

def _f1_score(y, pred):
    tp = np.sum((pred == 1) & (y == 1))
    fp = np.sum((pred == 1) & (y == 0))
    fn = np.sum((pred == 0) & (y == 1))
    if tp == 0:
        return 0.0
    p, r = tp / (tp + fp), tp / (tp + fn)
    return 2 * p * r / (p + r) if p + r > 0 else 0.0

def _roc_auc(y, prob):
    y = np.asarray(y)
    pos = y == 1
    n_pos, n_neg = int(pos.sum()), int((~pos).sum())
    if n_pos == 0 or n_neg == 0:
        return np.nan
    order = np.argsort(prob, kind="mergesort")
    ranks = np.empty(len(prob), dtype=float)
    ranks[order] = np.arange(1, len(prob) + 1)
    return (ranks[pos].sum() - n_pos * (n_pos + 1) / 2) / (n_pos * n_neg)

def _metrics(y, prob, thr=0.5):
    y = np.asarray(y)
    yi = y.astype(int)
    pred = (np.asarray(prob) >= thr).astype(int)
    acc = np.mean(pred == yi)
    ba = _balanced_acc(yi, pred)
    f1v = _f1_score(yi, pred)
    auc = _roc_auc(yi, prob)
    return acc, ba, f1v, auc

def _stem(x):
    s = str(x).rsplit("/", 1)[-1]
    while s.endswith(".wav"):
        s = s[:-4]
    return s.replace("audios-", "").replace("audios/", "")

def _sl(df, kv, common, meta_cols):
    return (df.loc[kv.isin(common)]
              .loc[:, ~df.columns.isin(meta_cols)]
              .set_index(kv[kv.isin(common)].values))

def _task(c):
    s = str(c)
    s = s.rsplit("/", 1)[-1]
    while s.endswith(".wav"):
        s = s[:-4]
    s = s.replace("audios-", "").replace("audios/", "")
    if s.startswith("PD_"):
        s = s[3:]
    elif s.startswith("HC_"):
        s = s[3:]
    s = re.sub(r"_\d+$", "", s)
    return s

def load_nv(task=None):
    hc = pd.read_csv(str(BASE) + '/data/neurovoz/metadata/metadata_hc.csv')
    pd_ = pd.read_csv(str(BASE) + '/data/neurovoz/metadata/metadata_pd.csv')
    mt = pd.concat([hc, pd_], ignore_index=True)
    mt["_pid"] = mt["ID"].map(lambda i: str(int(i)).zfill(4))
    mt["Age"] = pd.to_numeric(mt["Age"], errors="coerce")
    a_by_pid = mt.groupby("_pid")["Age"].median()
    mt["Sex"] = pd.to_numeric(mt["Sex"], errors="coerce")
    s_by_pid = mt.groupby("_pid")["Sex"].median()
    mt["UPDRS scale"] = pd.to_numeric(mt["UPDRS scale"], errors="coerce")
    u_by_pid = mt.groupby("_pid")["UPDRS scale"].median()

    grbas = {}
    for f in os.listdir(str(BASE) + '/data/grbas'):
        if not f.endswith(".csv"):
            continue
        for _, r in pd.read_csv(str(BASE) + '/data/grbas/' + f).iterrows():
            gid = os.path.splitext(os.path.basename(str(r["text_patient_id"])))[0].strip()
            try:
                grbas[gid] = float(r["G"])
            except (ValueError, TypeError):
                pass

    intens = pd.read_csv(str(NV_FEAT) + '/intensity_features_neurovoz.csv')
    pause = pd.read_csv(str(NV_FEAT) + '/pause_features_neurovoz.csv')
    syll = pd.read_csv(str(NV_FEAT) + '/syllable_features_neurovoz.csv')
    ki, kp, ks = (intens["tot_name"].apply(_stem),
                  pause["tot_name"].apply(_stem),
                  syll["tot_name"].apply(_stem))
    common = sorted(set(ki) & set(kp) & set(ks))
    if task is not None:
        common = [c for c in common if _task(c) == task]

    X = pd.concat([_sl(intens, ki, common, META_COLS),
                   _sl(pause, kp, common, META_COLS),
                   _sl(syll, ks, common, META_COLS)], axis=1)
    X = X.values.astype(np.float64)
    Y = np.array([1 if c.startswith("PD_") else 0 for c in common], dtype=np.float64)
    pid = np.array([c[-4:] for c in common])
    ag = np.array([a_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    sx = np.array([s_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    up = np.array([np.nan if np.isnan(u_by_pid.get(p, np.nan)) else
                   u_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    gr = np.array([grbas.get(c[3:], np.nan) for c in common], dtype=np.float64)
    ok = ~(np.isnan(X).any(axis=1) | np.isnan(ag) | np.isnan(sx))
    X, Y, ag, sx, pid, up, gr = X[ok], Y[ok], ag[ok], sx[ok], pid[ok], up[ok], gr[ok]
    if task is not None:
        return X, Y, ag, sx, pid, up, gr
    tasks = np.array([_task(c) for c in common])[ok]
    return X, Y, ag, sx, pid, up, gr, tasks

def load_pc():
    base = str(BASE) + '/data/parkceleb_extracted'
    intens = pd.read_csv(str(base) + '/intensity_features_parkceleb.csv')
    pause = pd.read_csv(str(base) + '/pause_features_parkceleb.csv')
    syll = pd.read_csv(str(base) + '/syllable_features_parkceleb.csv')
    kid = intens["audio_file"].map(
        lambda f: (lambda m: m.group(1) if m else np.nan)(
            age.IDRE.search(str(f))))
    m = kid.notna()
    intens, pause, syll, kid = (intens[m].reset_index(drop=True),
                                pause[m].reset_index(drop=True),
                                syll[m].reset_index(drop=True),
                                kid[m].reset_index(drop=True))
    X = pd.concat([intens.loc[:, ~intens.columns.isin(age.FEATURE_META)],
                   pause.loc[:, ~pause.columns.isin(age.FEATURE_META)],
                   syll.loc[:, ~syll.columns.isin(age.FEATURE_META)]], axis=1)
    X = X.values.astype(np.float64)
    Y = intens["parkceleb_label"].values.astype(np.float64)
    sid = intens["speaker_id"].astype(str).values

    ages = pd.read_csv(str(BASE) + '/data/parkceleb_ages.csv')
    age_by_af = (ages.drop_duplicates("audio_file").set_index("audio_file")["age"])
    af = intens["audio_file"].astype(str).map(
        lambda f: f if f in age_by_af.index else np.nan)
    t_age = af.map(age_by_af).values.astype(np.float64)

    upload = pd.read_csv(str(BASE) + '/data/parkceleb_upload_years.csv')
    upload["video_id"] = upload["video_id"].astype(str)
    up_full = upload["upload_date"].astype(str).str[:4]
    up = up_full[up_full.str.fullmatch(r"\d{4}")].astype(int)
    up.index = upload.loc[up.index, "video_id"]

    import openpyxl
    wb = openpyxl.load_workbook(str(BASE) + '/data/ParkCeleb/PD_demo.xlsx',
                                read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    hdr = [str(c).strip().lower().replace(" ", "") for c in rows[0]]
    yod = {}
    for r in rows[2:]:
        if r[0] is None:
            continue
        yod[str(r[0]).strip()] = int(str(r[hdr.index("yod")]).strip())
    wb2 = openpyxl.load_workbook(str(BASE) + '/data/ParkCeleb/speakers_pairs.xlsx',
                                 read_only=True)
    rows2 = list(wb2.active.iter_rows(values_only=True))
    pairs = {}
    for r in rows2[1:]:
        if r[0] is None:
            continue
        pairs[str(r[0]).strip()] = str(r[1]).strip()

    iv = np.full(len(X), np.nan)
    for i, sp in enumerate(sid):
        vid = kid[i]
        if vid not in up.index:
            continue
        uy = up[vid]
        yr = yod.get(sp) if sp.startswith("pd_") else yod.get(pairs.get(sp))
        if yr is None:
            continue
        iv[i] = uy - yr

    out = {"full": (X, Y, sid, t_age)}
    for name, (lo, hi) in INTERVALS.items():
        sel = (iv >= lo) & (iv <= hi)
        out[name] = (X[sel], Y[sel], sid[sel], t_age[sel])
    return out

def _gbc(X, y, w=None, seed=None):
    clf = GradientBoostingClassifier(n_estimators=100, learning_rate=0.1,
                                     max_depth=3,
                                     random_state=RANDOM_SEED if seed is None
                                     else seed)
    if w is not None:
        w = np.asarray(w, dtype=float)
        w = w * (len(X) / w.sum())
    clf.fit(np.asarray(X), y.astype(int), sample_weight=w)
    return clf

def _gbc_two_stage(X_base, y_base, w_base, X_spec, y_spec, w_spec,
                   extra_trees=20, fine_lr=0.1):
    clf = GradientBoostingClassifier(n_estimators=100, learning_rate=0.1,
                                     max_depth=3, random_state=RANDOM_SEED,
                                     warm_start=True)
    if w_base is not None:
        w_base = np.asarray(w_base, dtype=float)
        w_base = w_base * (len(X_base) / w_base.sum())
    clf.fit(np.asarray(X_base), y_base.astype(int), sample_weight=w_base)
    clf.n_estimators += extra_trees
    clf.learning_rate = fine_lr
    if w_spec is not None:
        w_spec = np.asarray(w_spec, dtype=float)
        w_spec = w_spec * (len(X_spec) / w_spec.sum())
    clf.fit(np.asarray(X_spec), y_spec.astype(int), sample_weight=w_spec)
    return clf

def _cb(X, y, ag, sx, model="age", up=None, gr=None):
    graph = GRAPH_MODELS[model]
    if "X->y" in graph and "y->X" not in graph:
        return np.ones((len(X), 2))
    fit = CB_FIT_METHOD
    kwargs = dict(causal_graph=graph, cause_var="y",
                  effect_var="X", fit_method=fit,
                  X=X, y=y, Age=(ag, N_AGE_BINS))
    if "Sex" in graph:
        if fit == "kde":
            jit = np.random.default_rng(len(X)).uniform(-0.005, 0.005, len(sx))
            sx = np.asarray(sx, dtype=float) + jit
        kwargs["Sex"] = (sx, 2)
    if "UPDRS" in graph:
        u = np.asarray(up, dtype=float)
        pd_mask = y == 1
        u_pd_med = np.nanmedian(u[pd_mask]) if np.any(np.isfinite(u[pd_mask])) else 0.0
        u = np.where(np.isnan(u), np.where(pd_mask, u_pd_med, 0.0), u)
        kwargs["UPDRS"] = (u, len(UPDRS_EDGES) - 1)
    if "GRBAS" in graph:
        g = np.asarray(gr, dtype=float)
        pd_mask = y == 1
        g_pd_med = np.nanmedian(g[pd_mask]) if np.any(np.isfinite(g[pd_mask])) else 0.0
        g = np.where(np.isnan(g), np.where(pd_mask, g_pd_med, 0.0), g)
        kwargs["GRBAS"] = (g, len(GRBAS_EDGES) - 1)
    return np.asarray(ca.causal_weights(**kwargs))

def _tcb_weights(U_src, U_trg, Y_trg, clip=None):
    U_src = np.asarray(U_src, dtype=float)
    U_trg = np.asarray(U_trg, dtype=float)
    if U_src.ndim == 1:
        U_src = U_src[:, None]
    if U_trg.ndim == 1:
        U_trg = U_trg[:, None]
    pu_src = np.asarray(gaussian_kde(U_src.T)(U_trg.T)).ravel()
    pu_trg = np.asarray(gaussian_kde(U_trg.T)(U_trg.T)).ravel()
    w = pu_src / pu_trg
    if clip is not None:
        lo, hi = np.percentile(w, clip)
        w = np.clip(w, lo, hi)
    return w / np.mean(w)

def _loc_weights(X, y, age):
    ag = np.asarray(age, dtype=float)
    med = np.nanmedian(ag) if np.any(np.isfinite(ag)) else 0.0
    ag = np.where(np.isfinite(ag), ag, med)
    W = np.asarray(ca.causal_weights(causal_graph=AGE_GRAPH, cause_var="y",
                                     effect_var="X", fit_method=CB_FIT_METHOD,
                                     X=X, y=y, Age=(ag, N_AGE_BINS)))
    return W

def _speaker_split(sid, y, seed):
    rng = np.random.default_rng(seed)
    uniq = np.unique(sid)
    maj = np.array([np.bincount(y[sid == u].astype(int)).argmax() for u in uniq])
    tr_u, te_u = [], []
    for c in (0, 1):
        idx = np.where(maj == c)[0]
        p = rng.permutation(idx)
        k = int(round(0.8 * len(p)))
        tr_u += list(uniq[p[:k]].tolist())
        te_u += list(uniq[p[k:]].tolist())
    tr = np.isin(sid, np.asarray(tr_u))
    te = np.isin(sid, np.asarray(te_u))
    return tr, te

def _perm_gt(a, b, n_flip=10000, seed=7):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    d = a - b
    d = d[np.isfinite(d)]
    if len(d) < 3 or np.std(d) == 0:
        return 1.0
    rng = np.random.default_rng(seed)
    obs = d.mean()
    flips = rng.integers(0, 2, size=(n_flip, len(d)), dtype=np.int8)
    null = (d[None, :] * (2 * flips - 1)).mean(axis=1)
    return float((null >= obs).mean())
