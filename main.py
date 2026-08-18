import argparse
import os

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import ks_2samp
from sklearn.decomposition import PCA
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import StratifiedKFold
from sklearn.svm import SVC
from tqdm import tqdm

import old.causal as c

OUT = str(c.BASE) + "/outputs"
QUALITY_CSV = str(c.BASE) + "/data/parkceleb_extracted/quality_features_parkceleb.csv"
N_FOLDS = 10
NV_EXCLUDED = {"0056", "0074", "0121", "0135"}
TCB_CLIP = (1.0, 99.0)
RANDOM_SEED = c.RANDOM_SEED

GRAPH_CONTROL = "y;X; y->X;"
GRAPH_AGE = "y;X;Age; y->X; Age->X; Age->y;"
GRAPH_ACCENT = "y;X;Accent; y->X; Accent->X; Accent->y;"
GRAPH_QUALITY = "y;X;Quality; y->X; Quality->X; Quality->y;"
GRAPH_SEX = "y;X;Sex; y->X; Sex->X; Sex->y;"

PC_COLS = ["variable", "n", "conf", "deconf", "dacc", "p_ks", "reps", "cb_repeats"]
NV_COLS = ["graph", "variant", "conf", "deconf", "dacc", "p_ks", "reps", "cb_repeats"]
NV1_COLS = ["graph", "conf", "deconf", "dacc", "p_ks", "reps", "cb_repeats"]
XFER_COLS = ["interval", "conf", "deconf", "dacc", "p_ks", "reps", "cb_repeats"]

NAT_CLUSTER = {
    "Canadian-America": 0, "British": 1, "English": 1, "Austrian-American": 2,
    "American-canadian": 0, "Irish": 2, "Italian": 2, "Spanish": 2, "Austrian": 2,
    "French": 2, "Romanian-born German": 2, "German": 2, "Scottish": 2,
    "Australian": 2, "Peruvian": 2, "Canadian": 0, "American": 0,
}


def done(path, key_cols, vals):
    try:
        d = pd.read_csv(path, dtype=str)
        return (pd.Index(key_cols).isin(d.columns).all()
                and tuple(map(str, vals)) in set(zip(*[d[k].values for k in key_cols])))
    except Exception:
        return False


def save(path, key_cols, cols, rows):
    df = pd.DataFrame(rows, columns=cols)
    if os.path.exists(path):
        try:
            df = pd.concat([pd.read_csv(path), df], ignore_index=True)
        except Exception:
            pass
    for k in key_cols:
        df[k] = df[k].astype(str)
    df.drop_duplicates(subset=key_cols, keep="last").to_csv(path, index=False)


def weights(graph, X, y, feats, fit="kde"):
    kw = dict(causal_graph=graph, cause_var="y", effect_var="X", fit_method=fit,
              X=X, y=y)
    for name, vals, bins, disc in feats:
        vals = np.asarray(vals, dtype=float)
        if disc and fit == "kde":
            rng = np.random.default_rng(abs(hash(name)) % (2 ** 32))
            vals = vals + rng.uniform(-0.05, 0.05, vals.size).reshape(vals.shape)
        kw[name] = (vals, bins)
    W = np.asarray(c.ca.causal_weights(**kw))
    return W if np.isfinite(W).all() else np.ones(np.shape(W))


def resample(X, y, W, pid, repeats, seed=None):
    rng = np.random.default_rng(RANDOM_SEED if seed is None else seed)
    X, y, W = np.asarray(X), np.asarray(y).astype(int), np.asarray(W, dtype=float)
    idx = []
    for d in (0, 1):
        m = y == d
        if not m.any():
            continue
        p = W[m, d] / W[m, d].sum()
        idx += [np.where(m)[0][rng.choice(int(m.sum()), int(m.sum()), p=p)]
                for _ in range(repeats)]
    idx = np.concatenate(idx)
    return X[idx], y[idx], np.asarray(pid)[idx]


def fold_masks(pid, y, seed):
    pid, y = np.asarray(pid), np.asarray(y).astype(int)
    uniq = np.unique(pid)
    maj = np.array([np.bincount(y[pid == u]).argmax() for u in uniq])
    skf = StratifiedKFold(n_splits=N_FOLDS, shuffle=True, random_state=seed)
    for tr, te in skf.split(uniq, maj):
        yield np.isin(pid, uniq[tr]), np.isin(pid, uniq[te])


def ks_p(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    a, b = a[np.isfinite(a)], b[np.isfinite(b)]
    if len(a) < 3 or len(b) < 3:
        return 1.0
    return float(ks_2samp(a, b, alternative="greater").pvalue)


def balacc(y, score):
    y = np.asarray(y).astype(int)
    return c._balanced_acc(y, (np.asarray(score) >= 0.5).astype(int))


def speaker_ba(sid, y, score):
    y = np.asarray(y).astype(int)
    sid = np.asarray(sid)
    score = np.asarray(score)
    yg, pg = [], []
    for s in np.unique(sid):
        m = sid == s
        yg.append(np.bincount(y[m]).argmax())
        pg.append(score[m].mean())
    return c._balanced_acc(np.array(yg), (np.array(pg) >= 0.5).astype(int))


def impute_median(v, y):
    v = np.asarray(v, dtype=float)
    y = np.asarray(y).astype(int)
    med = np.nanmedian(v[y == 1])
    med = med if np.isfinite(med) else 0.0
    return np.where(np.isnan(v), np.where(y == 1, med, 0.0), v)


def nv_data():
    X, y, ag, sx, pid, up, gr, _ = c.load_nv()
    m = ~np.isin(np.asarray(pid), list(NV_EXCLUDED))
    return X[m], y[m], ag[m], sx[m], pid[m], up[m], gr[m]


def nv_weights(graph, X, y, ag, sx, up, gr):
    if graph == GRAPH_CONTROL:
        return np.ones((len(X), 2))
    feats = [("Age", ag, 6, False), ("Sex", sx, 2, True)]
    if "UPDRS" in graph:
        feats.append(("UPDRS", impute_median(up, y), 5, False))
    if "GRBAS" in graph:
        feats.append(("GRBAS", impute_median(gr, y), 4, False))
    return weights(graph, X, y, feats)


def deconf_game(W, X, y, pid, reps, repeats, variant="weights", label=None):
    X, y, pid = np.asarray(X), np.asarray(y), np.asarray(pid)
    conf, dec = [], []
    for rep in tqdm(range(reps), desc=label, leave=False):
        for tr, te in fold_masks(pid, y, RANDOM_SEED + rep):
            seed = RANDOM_SEED + rep
            if variant == "weights":
                gd = c._gbc(X[tr], y[tr],
                            W[tr][np.arange(int(tr.sum())), y[tr].astype(int)],
                            seed=seed)
            else:
                Xd, yd, _ = resample(X[tr], y[tr], W[tr], pid[tr], repeats, seed=seed)
                gd = c._gbc(Xd, yd, seed=seed)
            gc = c._gbc(X[tr], y[tr], seed=seed)
            conf.append(balacc(y[te], gc.predict_proba(X[te])[:, 1]))
            dec.append(balacc(y[te], gd.predict_proba(X[te])[:, 1]))
    return np.asarray(conf), np.asarray(dec)


def conf_subsample(X, y, sex, p_m, p_f, seed):
    rng = np.random.default_rng(seed)
    y, sex = np.asarray(y).astype(int), np.asarray(sex).astype(int)
    keep = np.zeros(len(y), bool)
    for g, p in ((1, p_m), (0, p_f)):
        hc = (sex == g) & (y == 0)
        pd = (sex == g) & (y == 1)
        keep[hc] = True
        n_pd = min(int(round(int(hc.sum()) * p / max(1 - p, 1e-9))), int(pd.sum()))
        if n_pd:
            keep[np.where(pd)[0][rng.permutation(int(pd.sum()))[:n_pd]]] = True
    i = np.where(keep)[0]
    return np.asarray(X)[i], y[i], sex[i], i


def semisynth(X, y, sex, pid, repeats):
    Xc, yc, sc, ic = conf_subsample(X, y, sex, 0.8, 0.2, 101)
    Xr, yr, _, ir = conf_subsample(X, y, sex, 0.2, 0.8, 202)
    pidc, pidr = np.asarray(pid)[ic], np.asarray(pid)[ir]
    Wc = weights(GRAPH_SEX, Xc, yc, [("Sex", sc, 2, True)])
    Xd, yd, pidd = resample(Xc, yc, Wc, pidc, repeats)

    def loo(k, Xtr, ytr, pidtr, Xte, yte):
        ba = []
        for u in tqdm(np.unique(pidtr), leave=False):
            m = pidtr != u
            if len(np.unique(ytr[m])) < 2:
                continue
            g = c._gbc(Xtr[m], ytr[m], seed=RANDOM_SEED + k)
            ba.append(balacc(yte, g.predict_proba(Xte)[:, 1]))
        return np.mean(ba)

    return {"conf->conf": loo(1, Xc, yc, pidc, Xc, yc),
            "conf->rev": loo(2, Xc, yc, pidc, Xr, yr),
            "dec->conf": loo(3, Xd, yd, pidd, Xc, yc),
            "dec->rev": loo(4, Xd, yd, pidd, Xr, yr)}


def transfer(pc, xfer_reps, cb_repeats):
    Xp, yp, agp, sxp, _, up_p, gr_p = c.load_nv(task="PETACA_BLANCA")
    pid_pb = np.arange(len(yp))
    mu, sd = Xp.mean(0), np.maximum(Xp.std(0), 1e-9)
    Xs = (Xp - mu) / sd
    pca = PCA(n_components=2)
    L = pca.fit_transform(Xs[:, c.INTENSITY_IDX])
    corrs = np.array([abs(np.corrcoef(L[:, i], yp)[0, 1]) for i in range(2)])
    ku = np.argsort(corrs)[:c.U_DIMS]
    Unv = L[:, ku]
    path = OUT + "/pipeline_transfer.csv"
    rows = []
    for name, (X_pc, y_pc, sid_pc, _) in pc.items():
        if name == "full":
            continue
        if done(path, ["interval", "reps", "cb_repeats"],
                [name, xfer_reps, cb_repeats]):
            print("  interval %-3s already saved, skipping" % name)
            continue
        m0 = ~np.isnan(X_pc).any(axis=1)
        X_t = (X_pc[m0] - mu) / sd
        y_t, sid_t = y_pc[m0], sid_pc[m0]
        Upc = pca.transform(X_t[:, c.INTENSITY_IDX])[:, ku]
        Wt = np.column_stack([c._tcb_weights(Upc, Unv, yp, clip=TCB_CLIP)] * 2)
        med = {"conf": [], "dec": []}
        for rep in tqdm(range(xfer_reps), desc="interval %s" % name, leave=False):
            Xd, yd, _ = resample(Xs, yp, Wt, pid_pb, cb_repeats,
                                 seed=RANDOM_SEED + rep)
            gd = c._gbc(Xd, yd, seed=RANDOM_SEED + rep)
            gc = c._gbc(Xs, yp, seed=RANDOM_SEED + rep)
            med["dec"].append(speaker_ba(sid_t, y_t, gd.predict_proba(X_t)[:, 1]))
            med["conf"].append(speaker_ba(sid_t, y_t, gc.predict_proba(X_t)[:, 1]))
        mc, md = float(np.median(med["conf"])), float(np.median(med["dec"]))
        p = ks_p(med["dec"], med["conf"])
        print("  %-3s  conf=%.4f  deconf=%.4f  dacc=%+.4f  p_ks=%.4f" % (
            name, mc, md, md - mc, p))
        rows.append([name, mc, md, md - mc, p, xfer_reps, cb_repeats])
    if rows:
        save(path, ["interval"], XFER_COLS, rows)


def rep_pc():
    path = OUT + "/pipeline_replication_pc.csv"
    rows = []
    for name, (X_pc, y_pc, sid_pc, _) in c.load_pc().items():
        if name == "full" or len(np.unique(sid_pc)) < 8:
            continue
        if done(path, ["interval"], [name]):
            print("  interval %-3s already saved, skipping" % name)
            continue
        ba = []
        for seed in range(RANDOM_SEED, RANDOM_SEED + 5):
            for tr, te in fold_masks(sid_pc, y_pc, seed):
                sc = c._SS().fit(X_pc[tr])
                g = c._gbc(sc.transform(X_pc[tr]), y_pc[tr])
                ba.append(balacc(y_pc[te], g.predict_proba(sc.transform(X_pc[te]))[:, 1]))
        ba = float(np.mean(ba))
        print("  interval %-3s n=%-4d speakers=%-3d ba=%.4f" % (
            name, len(y_pc), len(np.unique(sid_pc)), ba))
        rows.append([name, len(y_pc), len(np.unique(sid_pc)), ba])
    if rows:
        save(path, ["interval"], ["interval", "n", "speakers", "ba"], rows)


def rep_nv():
    df = pd.read_csv(str(c.BASE) +
                     "/data/neurovoz/clean_data_featurized/full_data_processed_with_V.csv")
    df["user_id"] = df["user_id"].map(lambda s: str(s).zfill(4))
    df = df[~df["user_id"].isin(NV_EXCLUDED)]
    feats = [x for x in df.columns
             if x not in ("user_id", "word_id", "pk_label", "Sex", "Age")]
    ag = df.groupby("user_id")[feats + ["pk_label"]].median()
    X = ag[feats].values.astype(np.float64)
    y = ag["pk_label"].values.astype(int)
    pid = ag.index.values
    print("  patients=%d  pd=%d  hc=%d  features=%d" % (
        len(pid), int(y.sum()), len(y) - int(y.sum()), len(feats)))
    path = OUT + "/pipeline_replication_nv.csv"
    models = [("random forest", RandomForestClassifier(
        n_estimators=100, random_state=RANDOM_SEED)),
        ("logistic regression", LogisticRegression(
            l1_ratio=1, C=1.0, solver="saga", random_state=RANDOM_SEED,
            max_iter=10000))]
    rows = []
    for name, model in models:
        if done(path, ["model"], [name]):
            print("  %-20s  already saved, skipping" % name)
            continue
        ba = []
        for seed in range(RANDOM_SEED, RANDOM_SEED + 5):
            for tr, te in fold_masks(pid, y, seed):
                sc = c._SS().fit(X[tr])
                model.fit(sc.transform(X[tr]), y[tr])
                ba.append(balacc(y[te], model.predict_proba(sc.transform(X[te]))[:, 1]))
        ba = float(np.mean(ba))
        print("  %-20s  ba=%.4f" % (name, ba))
        rows.append([name, ba])
    if rows:
        save(path, ["model"], ["model", "ba"], rows)


def accent_map():
    acc = {}
    for path in (str(c.BASE) + "/data/ParkCeleb/PD_demo.xlsx",
                 str(c.BASE) + "/data/ParkCeleb/CN_demo.xlsx"):
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        hdr = [str(x).strip().lower().replace(" ", "") for x in rows[0]]
        i_s, i_n, i_e = (hdr.index("speakerid"), hdr.index("nationality"),
                         hdr.index("ethnicity"))
        for r in rows[1:]:
            if r[i_s] is None:
                continue
            sid = str(r[i_s]).strip().lower()
            nat = str(r[i_n]).strip() if r[i_n] is not None else ""
            eth = str(r[i_e]).strip() if r[i_e] is not None else ""
            if nat not in NAT_CLUSTER:
                continue
            a = NAT_CLUSTER[nat]
            if eth and eth.lower() != "white":
                a += 2
            acc[sid] = int(a)
    return acc


def quality_arms(path4, rows, Xf, yf, sidf, mq, qv, qcols, pc_reps, cb_repeats):
    def run(var, label, vals, bins):
        if done(path4, ["variable", "reps", "cb_repeats"],
                [var, pc_reps, cb_repeats]):
            print("  %-36s already saved, skipping" % label)
            return
        W = weights(GRAPH_QUALITY, Xf[mq], yf[mq], [("Quality", vals, bins, False)])
        conf, dec = deconf_game(W, Xf[mq], yf[mq], sidf[mq], pc_reps, cb_repeats,
                                label=label)
        mc, md = float(np.median(conf)), float(np.median(dec))
        p = ks_p(dec, conf)
        print("  %-36s n=%-5d conf=%.4f  deconf=%.4f  dacc=%+.4f  p_ks=%.4f" % (
            label, len(vals), mc, md, md - mc, p))
        rows.append([var, len(vals), mc, md, md - mc, p, pc_reps, cb_repeats])

    run("quality", "pc quality", qv[mq], [5] * len(qcols))
    sco = []
    for j in range(len(qcols)):
        ba = []
        for tr, te in fold_masks(sidf[mq], yf[mq], RANDOM_SEED):
            svc = SVC(C=1.0).fit(qv[mq][tr, j:j + 1], yf[mq].astype(int)[tr])
            ba.append(balacc(yf[mq].astype(int)[te], svc.predict(qv[mq][te, j:j + 1])))
        sco.append(float(np.mean(ba)))
    jb = int(np.argmax(sco))
    run("quality-" + qcols[jb], "pc quality-" + qcols[jb],
        qv[mq][:, jb:jb + 1], [10])


def main():
    global RANDOM_SEED
    ap = argparse.ArgumentParser()
    ap.add_argument("--reps", type=int, default=None)
    ap.add_argument("--pc-reps", type=int, default=50)
    ap.add_argument("--nv-reps", type=int, default=200)
    ap.add_argument("--xfer-reps", type=int, default=200)
    ap.add_argument("--cb-repeats", type=int, default=10)
    ap.add_argument("--seed", type=int, default=c.RANDOM_SEED)
    args = ap.parse_args()
    RANDOM_SEED = args.seed
    if args.reps is not None:
        args.pc_reps = args.nv_reps = args.xfer_reps = args.reps
    os.makedirs(OUT, exist_ok=True)

    rep_pc()
    rep_nv()

    acc = accent_map()
    pc = c.load_pc()
    Xf, yf, sidf, agef = pc["full"]
    accent_all = np.array([acc.get(s, np.nan) for s in sidf], dtype=float)
    ok_a = ~np.isnan(accent_all)
    q = pd.read_csv(QUALITY_CSV) if os.path.exists(QUALITY_CSV) else None

    X, y, ag, sx, pid, up, gr = nv_data()
    print("  neurovoz: Age pd=%.1f hc=%.1f | Sex pd=%.1f hc=%.1f | UPDRS pd=%.1f | GRBAS pd=%.2f hc=%.2f" % (
        ag[y.astype(int) == 1].mean(), ag[y.astype(int) == 0].mean(),
        sx[y.astype(int) == 1].mean(), sx[y.astype(int) == 0].mean(),
        np.nanmedian(up[y.astype(int) == 1]),
        np.nanmean(gr[y.astype(int) == 1]), np.nanmean(gr[y.astype(int) == 0])))

    path3 = OUT + "/pipeline_variants_nv.csv"
    rows3 = []
    for graph, variant in (("control", "bootstrap"), ("control", "weights"),
                           ("ug-as[a]", "bootstrap"), ("ug-as[a]", "weights")):
        if done(path3, ["graph", "variant", "reps", "cb_repeats"],
                [graph, variant, args.nv_reps, args.cb_repeats]):
            print("  nv %-8s %-9s already saved, skipping" % (graph, variant))
            continue
        gx = GRAPH_CONTROL if graph == "control" else c.GRAPH_MODELS[graph]
        W = nv_weights(gx, X, y, ag, sx, up, gr)
        conf, dec = deconf_game(W, X, y, pid, args.nv_reps, args.cb_repeats, variant,
                                label="nv %s %s" % (graph, variant))
        mc, md = float(np.median(conf)), float(np.median(dec))
        p = ks_p(dec, conf)
        print("  nv %-8s %-9s conf=%.4f  deconf=%.4f  dacc=%+.4f  p_ks=%.4f" % (
            graph, variant, mc, md, md - mc, p))
        rows3.append([graph, variant, mc, md, md - mc, p,
                      args.nv_reps, args.cb_repeats])
    if rows3:
        save(path3, ["graph", "variant"], NV_COLS, rows3)

    path4 = OUT + "/pipeline_samecorpus_pc.csv"
    rows4 = []

    def pc_arm(var, graph, feat, X, y, sid, vals, bins, disc, label, n):
        if done(path4, ["variable", "reps", "cb_repeats"],
                [var, args.pc_reps, args.cb_repeats]):
            print("  %-14s already saved, skipping" % label)
            return
        W = weights(graph, X, y, [(feat, vals, bins, disc)])
        conf, dec = deconf_game(W, X, y, sid, args.pc_reps, args.cb_repeats, label=label)
        mc, md = float(np.median(conf)), float(np.median(dec))
        p = ks_p(dec, conf)
        print("  %-14s n=%-5d conf=%.4f  deconf=%.4f  dacc=%+.4f  p_ks=%.4f" % (
            label, n, mc, md, md - mc, p))
        rows4.append([var, n, mc, md, md - mc, p, args.pc_reps, args.cb_repeats])

    m_age = ~np.isnan(agef)
    if len(np.unique(sidf[m_age])) >= 8:
        pc_arm("age", GRAPH_AGE, "Age", Xf[m_age], yf[m_age], sidf[m_age],
               agef[m_age], 6, False, "pc age", int(m_age.sum()))
    if ok_a.any() and len(np.unique(sidf[ok_a])) >= 8:
        pc_arm("accent", GRAPH_ACCENT, "Accent", Xf[ok_a], yf[ok_a], sidf[ok_a],
               accent_all[ok_a], len(np.unique(accent_all[ok_a])), True,
               "pc accent", int(ok_a.sum()))
    if q is not None:
        qcols = [x for x in q.columns
                 if x not in ("video_id", "audio_file", "speaker_id")]
        qd = {str(k): np.asarray(v, dtype=float)
              for k, v in zip(q["video_id"].astype(str), q[qcols].values)}
        intens = pd.read_csv(str(c.BASE) +
                             "/data/parkceleb_extracted/intensity_features_parkceleb.csv")
        kid = intens["audio_file"].map(
            lambda f: (lambda mm: mm.group(1) if mm else np.nan)(
                c.age.IDRE.search(str(f))))
        mk = kid.notna().values
        kidr = kid[mk].values
        sidr = intens.loc[mk, "speaker_id"].astype(str).values
        upload = pd.read_csv(str(c.BASE) + "/data/parkceleb_upload_years.csv")
        upf = upload["upload_date"].astype(str).str[:4]
        yok = upf.str.fullmatch(r"\d{4}")
        uy = {str(v): int(y_) for v, y_ in zip(upload.loc[yok, "video_id"], upf[yok])}
        wb = openpyxl.load_workbook(str(c.BASE) + "/data/ParkCeleb/PD_demo.xlsx",
                                    read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        hdr = [str(x).strip().lower().replace(" ", "") for x in rows[0]]
        yod = {str(r[0]).strip(): int(str(r[hdr.index("yod")]).strip())
               for r in rows[2:] if r[0] is not None}
        wb2 = openpyxl.load_workbook(str(c.BASE) +
                                     "/data/ParkCeleb/speakers_pairs.xlsx",
                                     read_only=True)
        rows2 = list(wb2.active.iter_rows(values_only=True))
        pairs = {str(r[0]).strip(): str(r[1]).strip()
                 for r in rows2[1:] if r[0] is not None}
        iv = np.full(len(kidr), np.nan)
        for i, (sp, vid) in enumerate(zip(sidr, kidr)):
            if vid not in uy:
                continue
            yr = yod.get(sp) if sp.startswith("pd_") else yod.get(pairs.get(sp))
            if yr is None:
                continue
            iv[i] = uy[vid] - yr
        cent = np.zeros(len(kidr), bool)
        for lo, hi in c.INTERVALS.values():
            cent |= (iv >= lo) & (iv <= hi)
        qv = np.array([qd.get(str(v), np.full(len(qcols), np.nan)) for v in kidr])
        mq = cent & ~np.isnan(qv).any(axis=1)
        quality_arms(path4, rows4, Xf, yf, sidf, mq, qv, qcols,
                     args.pc_reps, args.cb_repeats)
    elif q is None and not done(path4, ["variable", "reps", "cb_repeats"],
                                ["quality", args.pc_reps, args.cb_repeats]):
        rows4.append(["quality", 0, np.nan, np.nan, np.nan, np.nan,
                      args.pc_reps, args.cb_repeats])
    if rows4:
        save(path4, ["variable"], PC_COLS, rows4)

    path4b = OUT + "/pipeline_samecorpus_nv.csv"
    if done(path4b, ["graph", "reps", "cb_repeats"],
            ["g-as[a]", args.nv_reps, args.cb_repeats]):
        print("  nv g-as[a]  already saved, skipping")
    else:
        W = nv_weights(c.GRAPH_MODELS["g-as[a]"], X, y, ag, sx, up, gr)
        conf, dec = deconf_game(W, X, y, pid, args.nv_reps, args.cb_repeats, "bootstrap",
                                label="nv g-as[a]")
        mc, md = float(np.median(conf)), float(np.median(dec))
        p = ks_p(dec, conf)
        print("  nv g-as[a]  conf=%.4f  deconf=%.4f  dacc=%+.4f  p_ks=%.4f" % (
            mc, md, md - mc, p))
        save(path4b, ["graph"], NV1_COLS,
             [["g-as[a]", mc, md, md - mc, p, args.nv_reps, args.cb_repeats]])

    path4c = OUT + "/pipeline_semisynth.csv"
    cases = ("conf->conf", "conf->rev", "dec->conf", "dec->rev")
    if all(done(path4c, ["case", "cb_repeats"], [case, args.cb_repeats])
           for case in cases):
        print("  semisynth  all cases already saved, skipping")
    else:
        res = semisynth(X, y, sx, pid, args.cb_repeats)
        save(path4c, ["case"], ["case", "ba", "cb_repeats"],
             [[case, res[case], args.cb_repeats] for case in cases])

    transfer(pc, args.xfer_reps, args.cb_repeats)


if __name__ == "__main__":
    main()