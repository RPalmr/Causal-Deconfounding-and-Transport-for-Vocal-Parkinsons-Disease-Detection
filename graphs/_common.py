import sys, re, os
import numpy as np
import pandas as pd
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
for p in [str(BASE) + '/libraries/CausalBootstrapping', str(BASE) + '/libraries/GRAPL',
          str(BASE) + '/libraries/causal_assistant', str(BASE) + '/tCB', str(BASE) + '/tests']:
    if p not in sys.path:
        sys.path.insert(0, p)

import causal_assistant as ca
import sID_expr as sid_alg
import grapl.dsl as dsl
from scipy.stats import gaussian_kde, wilcoxon
from sklearn.ensemble import GradientBoostingClassifier

import transport_u_age as age

RANDOM_SEED = 42
N_BOOT = 50
N_AGE_BINS = 6
CB_FIT_METHOD = "kde"

CORPUS_SEP_MAX = 1.0
META_COLS = {"Unnamed: 0", "tot_name"}
INTERVALS = {"-2": (-10, -6), "-1": (-5, -1), "1": (0, 5), "2": (6, 10)}
NV_FEAT = str(BASE) + '/data/neurovoz/extracted_features'
INTENSITY_IDX = slice(0, 38)
N_FEAT = 66

AGE_GRAPH = "y;X;Age; y->X; Age->X; Age->y;"
SEL_DIAGRAM = "y;X;Age;U;S; y->X; U->X; Age->X; Age->y; S->U;"
UPDRS_EDGES = [0, 10, 20, 40, 70, 200]
GRBAS_EDGES = [0, 1, 2, 3, 4]
GRAPH_MODELS = {
    "age":   "y;X;Age; y->X; Age->X; Age->y;",
    "grbas": "y;X;Age;GRBAS; y->X; Age->X; Age->y; y->GRBAS; GRBAS->X;",
    "updrs": "y;X;Age;UPDRS; y->X; Age->X; Age->y; y->UPDRS; UPDRS->X;",
    "both":  "y;X;Age;UPDRS;GRBAS; y->X; Age->X; Age->y; y->UPDRS; UPDRS->X; "
             "y->GRBAS; GRBAS->X;",
}

def _zscore_fit(X):
    mu = X.mean(0)
    sd = X.std(0)
    return (X - mu) / np.maximum(sd, 1e-12), (mu, sd)

def _zscore_apply(X, p):
    mu, sd = p
    return (X - mu) / np.maximum(sd, 1e-12)

def _pca_fit(X, n_components):
    mu = X.mean(0)
    Xc = X - mu
    _, _, Vt = np.linalg.svd(Xc, full_matrices=False)
    Vt = _svd_flip(Vt)
    return mu, Vt[:n_components]

def _pca_apply(X, p):
    mu, comps = p
    return (X - mu) @ comps.T

def _svd_flip(Vt):
    Vt = Vt.copy()
    mx = np.argmax(np.abs(Vt), axis=1)
    signs = np.sign(Vt[np.arange(Vt.shape[0]), mx])
    signs[signs == 0] = 1.0
    return Vt * signs[:, None]

def _balanced_acc(y, pred):
    return 0.5 * (np.mean(pred[y == 0] == 0) + np.mean(pred[y == 1] == 1))

def _f1_score(y, pred):
    tp = np.sum((pred == 1) & (y == 1))
    fp = np.sum((pred == 1) & (y == 0))
    fn = np.sum((pred == 0) & (y == 1))
    if tp == 0:
        return 0.0
    p, r = tp / (tp + fp), tp / (tp + fn)
    return 2 * p * r / (p + r) if p + r > 0 else 0.0

def _roc_auc(y, prob):
    y = np.asarray(y)
    pos = y == 1
    n_pos, n_neg = int(pos.sum()), int((~pos).sum())
    if n_pos == 0 or n_neg == 0:
        return np.nan
    order = np.argsort(prob, kind="mergesort")
    ranks = np.empty(len(prob), dtype=float)
    ranks[order] = np.arange(1, len(prob) + 1)
    return (ranks[pos].sum() - n_pos * (n_pos + 1) / 2) / (n_pos * n_neg)

def _corp_sep(a, b):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    a, b = a[np.isfinite(a)], b[np.isfinite(b)]
    if len(a) < 2 or len(b) < 2:
        return 0.5
    val = np.concatenate([a, b])
    lab = np.concatenate([np.zeros(len(a)), np.ones(len(b))])
    return abs(_roc_auc(lab, val) - 0.5)

def _strat_folds(y, n_splits=3, seed=RANDOM_SEED):
    y = np.asarray(y).astype(int)
    rng = np.random.RandomState(seed)
    folds = [[] for _ in range(n_splits)]
    for c in np.unique(y):
        idx = rng.permutation(np.where(y == c)[0])
        for k, i in enumerate(idx):
            folds[k % n_splits].append(i)
    all_idx = np.arange(len(y))
    out = []
    for f in folds:
        te = np.asarray(f, dtype=int)
        tr = np.setdiff1d(all_idx, te)
        out.append((tr, te))
    return out

def _stem(x):
    s = str(x).rsplit("/", 1)[-1]
    while s.endswith(".wav"):
        s = s[:-4]
    return s.replace("audios-", "").replace("audios/", "")

def _sl(df, kv, common, meta_cols):
    return (df.loc[kv.isin(common)]
              .loc[:, ~df.columns.isin(meta_cols)]
              .set_index(kv[kv.isin(common)].values))

def _rmse(v):
    return np.sqrt(np.mean(np.asarray(v, dtype=float) ** 2))

def _sign_agr(a, b):
    a, b = np.asarray(a), np.asarray(b)
    m = (np.abs(a) + np.abs(b)) > 0
    if m.sum() == 0:
        return 1.0
    return float(np.mean(np.sign(a[m]) == np.sign(b[m])))

def _nv_col_order():
    cols = []
    for f in ("intensity_features_neurovoz.csv", "pause_features_neurovoz.csv",
              "syllable_features_neurovoz.csv"):
        cols += [c for c in pd.read_csv(str(NV_FEAT) + '/' + str(f), nrows=0)
                 if c not in META_COLS]
    return cols

def load_nv():
    hc = pd.read_csv(str(BASE) + '/data/neurovoz/metadata/metadata_hc.csv')
    pd_ = pd.read_csv(str(BASE) + '/data/neurovoz/metadata/metadata_pd.csv')
    mt = pd.concat([hc, pd_], ignore_index=True)
    mt["_pid"] = mt["ID"].map(lambda i: f"{int(i):04d}")
    mt["Age"] = pd.to_numeric(mt["Age"], errors="coerce")
    a_by_pid = mt.groupby("_pid")["Age"].median()
    mt["UPDRS scale"] = pd.to_numeric(mt["UPDRS scale"], errors="coerce")
    u_by_pid = mt.groupby("_pid")["UPDRS scale"].median()
    mt["Time Disease (years)"] = pd.to_numeric(mt["Time Disease (years)"],
                                               errors="coerce")
    d_by_pid = mt.groupby("_pid")["Time Disease (years)"].median()

    grbas = {}
    for f in os.listdir(str(BASE) + '/data/grbas'):
        if not f.endswith(".csv"):
            continue
        for _, r in pd.read_csv(str(BASE) + '/data/grbas/' + f).iterrows():
            gid = os.path.splitext(os.path.basename(str(r["text_patient_id"])))[0].strip()
            try:
                grbas[gid] = float(r["G"])
            except (ValueError, TypeError):
                pass

    intens = pd.read_csv(str(NV_FEAT) + '/intensity_features_neurovoz.csv')
    pause = pd.read_csv(str(NV_FEAT) + '/pause_features_neurovoz.csv')
    syll = pd.read_csv(str(NV_FEAT) + '/syllable_features_neurovoz.csv')
    ki, kp, ks = (intens["tot_name"].apply(_stem),
                  pause["tot_name"].apply(_stem),
                  syll["tot_name"].apply(_stem))
    common = sorted(set(ki) & set(kp) & set(ks))
    X = pd.concat([_sl(intens, ki, common, META_COLS),
                   _sl(pause, kp, common, META_COLS),
                   _sl(syll, ks, common, META_COLS)], axis=1)
    X = X.values.astype(np.float64)
    Y = np.array([1 if c.startswith("PD_") else 0 for c in common], dtype=np.float64)
    pid = np.array([c[-4:] for c in common])
    ag = np.array([a_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    up = np.array([np.nan if np.isnan(u_by_pid.get(p, np.nan)) else
                   u_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    gr = np.array([grbas.get(c[3:], np.nan) for c in common], dtype=np.float64)
    dr = np.array([d_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    ok = ~(np.isnan(X).any(axis=1) | np.isnan(ag))
    return X[ok], Y[ok], ag[ok], pid[ok], up[ok], gr[ok], dr[ok]

def load_pc():
    base = str(BASE) + '/data/parkceleb_extracted'
    intens = pd.read_csv(str(base) + '/intensity_features_parkceleb.csv')
    pause = pd.read_csv(str(base) + '/pause_features_parkceleb.csv')
    syll = pd.read_csv(str(base) + '/syllable_features_parkceleb.csv')
    kid = intens["audio_file"].map(
        lambda f: (lambda m: m.group(1) if m else np.nan)(
            age.IDRE.search(str(f))))
    m = kid.notna()
    intens, pause, syll, kid = (intens[m].reset_index(drop=True),
                                pause[m].reset_index(drop=True),
                                syll[m].reset_index(drop=True),
                                kid[m].reset_index(drop=True))
    X = pd.concat([intens.loc[:, ~intens.columns.isin(age.FEATURE_META)],
                   pause.loc[:, ~pause.columns.isin(age.FEATURE_META)],
                   syll.loc[:, ~syll.columns.isin(age.FEATURE_META)]], axis=1)
    X = X.values.astype(np.float64)
    Y = intens["parkceleb_label"].values.astype(np.float64)
    sid = intens["speaker_id"].astype(str).values

    ages = pd.read_csv(str(BASE) + '/data/parkceleb_ages.csv')
    age_by_af = (ages.drop_duplicates("audio_file").set_index("audio_file")["age"])
    af = intens["audio_file"].astype(str).map(
        lambda f: f if f in age_by_af.index else np.nan)
    t_age = af.map(age_by_af).values.astype(np.float64)

    upload = pd.read_csv(str(BASE) + '/data/parkceleb_upload_years.csv')
    upload["video_id"] = upload["video_id"].astype(str)
    up_full = upload["upload_date"].astype(str).str[:4]
    up = up_full[up_full.str.fullmatch(r"\d{4}")].astype(int)
    up.index = upload.loc[up.index, "video_id"]

    import openpyxl
    wb = openpyxl.load_workbook(str(BASE) + '/data/ParkCeleb/PD_demo.xlsx',
                                read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    hdr = [str(c).strip().lower().replace(" ", "") for c in rows[0]]
    yod = {}
    for r in rows[2:]:
        if r[0] is None:
            continue
        yod[str(r[0]).strip()] = int(str(r[hdr.index("yod")]).strip())
    wb2 = openpyxl.load_workbook(str(BASE) + '/data/ParkCeleb/speakers_pairs.xlsx',
                                 read_only=True)
    rows2 = list(wb2.active.iter_rows(values_only=True))
    pairs = {}
    for r in rows2[1:]:
        if r[0] is None:
            continue
        pairs[str(r[0]).strip()] = str(r[1]).strip()

    iv = np.full(len(X), np.nan)
    for i, sp in enumerate(sid):
        vid = kid[i]
        if vid not in up.index:
            continue
        uy = up[vid]
        yr = yod.get(sp) if sp.startswith("pd_") else yod.get(pairs.get(sp))
        if yr is None:
            continue
        iv[i] = uy - yr

    out = {"full": (X, Y, sid, t_age, iv)}
    for name, (lo, hi) in INTERVALS.items():
        sel = (iv >= lo) & (iv <= hi)
        out[name] = (X[sel], Y[sel], sid[sel], t_age[sel], iv[sel])
    return out

def load_italian():
    return None, None, None, None


def _u_axis(X_nv, X_tgt_all, y_nv):
    nv_i = X_nv[:, INTENSITY_IDX]
    stacked = np.vstack([nv_i, X_tgt_all[:, INTENSITY_IDX]])
    zs = np.maximum(1e-12, stacked.std(0))
    ns = (nv_i - stacked.mean(0)) / zs
    mu2, comps = _pca_fit(ns, 5)
    Z = _pca_apply(ns, (mu2, comps))
    corrs = np.array([abs(np.corrcoef(Z[:, k], y_nv)[0, 1]) for k in range(5)])
    k = int(np.argmin(corrs))
    axis = (stacked.mean(0), zs, mu2, comps, k)
    return Z[:, k], axis

def _u_proj(X_t, axis):
    stacked_mean, zs, mu2, comps, k = axis
    ns = (X_t[:, INTENSITY_IDX] - stacked_mean) / zs
    return _pca_apply(ns, (mu2, comps))[:, k]

def _ratio(U_s, U_t):
    U_s = np.asarray(U_s, dtype=float)
    U_t = np.asarray(U_t, dtype=float)
    U_t = U_t[np.isfinite(U_t)]
    if len(U_t) < 2 or np.std(U_t) < 1e-9 or np.std(U_s) < 1e-9:
        return np.ones(len(U_s))
    k_t, k_s = gaussian_kde(U_t), gaussian_kde(U_s)
    r = k_t(U_s) / np.maximum(k_s(U_s), 1e-12)
    lo, hi = np.percentile(r, 1), np.percentile(r, 99)
    return np.clip(r, lo, hi)

def _cb(X, y, ag, model="age", up=None, gr=None, z=None):
    graph = GRAPH_MODELS[model] if model in GRAPH_MODELS else model
    kwargs = dict(causal_graph=graph, cause_var="y",
                  effect_var="X", fit_method=CB_FIT_METHOD,
                  X=X, y=y)
    if "Age" in graph:
        kwargs["Age"] = (ag, N_AGE_BINS)
    if "Z" in graph:
        if z is None:
            raise ValueError("graph contains Z but no z values passed")
        zv = np.asarray(z, dtype=float)
        zv = np.where(np.isnan(zv), np.nanmedian(zv), zv)
        kwargs["Z"] = (zv, N_AGE_BINS)
    if "UPDRS" in graph:
        u = np.asarray(up, dtype=float)
        u = np.where(np.isnan(u), 0.0, u)
        kwargs["UPDRS"] = (u, len(UPDRS_EDGES) - 1)
    if "GRBAS" in graph:
        g = np.asarray(gr, dtype=float)
        g = np.where(np.isnan(g), 0.0, g)
        kwargs["GRBAS"] = (g, len(GRBAS_EDGES) - 1)
    W = np.asarray(ca.causal_weights(**kwargs))
    return W

def _effect_vec(X, w1, w0):
    s1, s0 = w1.sum(), w0.sum()
    if s1 <= 0 or s0 <= 0:
        return np.zeros(X.shape[1])
    return (X * w1[:, None]).sum(0) / s1 - (X * w0[:, None]).sum(0) / s0

def _local_effect(Xt, yt, age_t):
    med = np.nanmedian(age_t)
    a = np.where(np.isfinite(age_t), age_t, med) - med
    mu, sd = Xt.mean(0), Xt.std(0)
    sd = np.maximum(sd, np.median(sd) * 0.05)
    Xz = (Xt - mu) / sd
    A = np.column_stack([np.ones(len(yt)), yt, a])
    return np.linalg.lstsq(A, Xz, rcond=None)[0][1]

def _cluster_sel(units, rng):
    u = np.unique(units)
    return np.isin(units, rng.choice(u, size=len(u), replace=True))

def _run_target(Xs, ys, ags, pids, Xt, yt, sidt, agt, r, rng, n_boot,
                model="age", ups=None, grs=None, z=None):
    cb, raw, loc = [], [], []
    for _ in range(n_boot):
        s = _cluster_sel(pids, rng)
        us_ = None if ups is None else ups[s]
        gs_ = None if grs is None else grs[s]
        zs_ = None if z is None else z[s]
        W = _cb(Xs[s], ys[s], ags[s], model=model, up=us_, gr=gs_, z=zs_)
        cb.append(_effect_vec(Xs[s], W[:, 1] * r[s], W[:, 0] * r[s]))
        raw.append(_effect_vec(Xs[s], (ys[s] == 1).astype(float),
                               (ys[s] == 0).astype(float)))
        th = _cluster_sel(sidt, rng)
        loc.append(_local_effect(Xt[th], yt[th], agt[th]))
    return dict(cb=np.array(cb), raw=np.array(raw),
                loc=np.array(loc))

GBC_PARAMS = dict(n_estimators=100, learning_rate=0.1, max_depth=3,
                  random_state=RANDOM_SEED)
CLF_CV_FOLDS = 3

def _gbc(X, y, w=None):
    clf = GradientBoostingClassifier(**GBC_PARAMS)
    if w is not None:
        w = np.asarray(w, dtype=float)
        w = w * (len(X) / w.sum())
    clf.fit(np.asarray(X), y.astype(int), sample_weight=w)
    return clf

def _metrics(y, prob):
    y = np.asarray(y)
    yi = y.astype(int)
    pred = (np.asarray(prob) >= 0.5).astype(int)
    acc = np.mean(pred == yi)
    ba = _balanced_acc(yi, pred)
    f1v = _f1_score(yi, pred)
    auc = _roc_auc(yi, prob)
    return acc, ba, f1v, auc

def _fmt(m):
    return str(round((m[0]), 2)) + '/' + str(round((m[1]), 2)) + '/' + str(round((m[2]), 2)) + '/' + str(round((m[3]), 2))

def _source_cv(Xs, ys, ws, folds=CLF_CV_FOLDS):
    folds_idx = _strat_folds(ys, folds, RANDOM_SEED)
    out = {}
    for arm, w in (("raw", None), ("cb", ws)):
        acc = []
        for tr, te in folds_idx:
            g = _gbc(Xs[tr], ys[tr], None if w is None else w[tr])
            acc.append(_metrics(ys[te], g.predict_proba(Xs[te])[:, 1]))
        out[arm] = np.mean(np.asarray(acc), axis=0)
    return out

FIG1 = {
    "a": dict(
        name="Figure 1(a) -- Z=age common cause of y and X; recalibrate P*(Z)/P(Z)",
        sel="y;X;Z;S; y->X; Z->X; Z->y; y<->X; S->Z;",
        graph="y;X;Z; y->X; Z->X; Z->y;",
        formula="P*(X|do(y)) = \\sum_Z P(X|y,Z) P*(Z)",
    ),
    "b": dict(
        name="Figure 1(b) -- Z=proxy, no causal path to X or y; directly transportable",
        sel="y;X;Z;S; y->X; y<->X; S->Z;",
        graph="y;X; y->X;",
        formula="P*(X|do(y)) = P(X|do(y))   [no recalibration]",
    ),
    "c": dict(
        name="Figure 1(c) -- Z=mediator between y and X; recalibrate P*(Z|y)",
        sel="y;X;Z;S; y->X; y->Z; Z->X; y<->X; S->Z;",
        graph="y;X;Z; y->X; y->Z; Z->X;",
        formula="P*(X|do(y)) = \\sum_Z P(X|do(y),Z) P*(Z|y)",
    ),
}

def main(fig="a"):
    cfg = FIG1[fig]
    print("sID transportability check ...", flush=True)
    try:
        g = dsl.GraplDSL().readgrapl(cfg["sel"])
        f = sid_alg.sid(Y={"X"}, X={"y"}, S={"S"}, D=g)
        print("   " + cfg["formula"], flush=True)
        print("   sID: " + (None if f is None else f.rhs.tocondstr()),
              flush=True)
    except Exception as e:
        print("   sID skipped:", type(e).__name__, e, flush=True)
    print()
    print("== " + cfg["name"] + " ==", flush=True)

    Xs, ys, ags, pids, _, _, drs_nv = load_nv()
    if ags.size > 0:
        age_str = 'age [' + str(int(np.nanmin(ags))) + '-' + str(int(np.nanmax(ags))) + ']'
    else:
        age_str = 'age [no data]'
    print('nv: ' + str(Xs.shape[0]) + ' recordings (' + str(int(ys.sum())) + ' pd / ' + str(len(ys) - int(ys.sum())) + ' hc), ' + age_str, flush=True)

    mu_s, sd_s = Xs.mean(0), Xs.std(0)
    sd_s = np.maximum(sd_s, np.median(sd_s) * 0.05)
    Xs = (Xs - mu_s) / sd_s

    pc = load_pc()
    zt = lambda X: (np.asarray(X, dtype=float) - mu_s) / sd_s
    targets = {}
    for name, (X, y, s, a, iv) in pc.items():
        m = ~np.isnan(X).any(axis=1)
        targets[name.upper()] = (zt(X[m]), y[m], s[m], a[m], iv[m])

    Zfull = Xs
    pool = np.vstack([zt(pc[k][0]) for k in pc])
    sep = np.array([_corp_sep(Zfull[:, f], pool[:, f]) for f in range(N_FEAT)])
    keep = sep <= CORPUS_SEP_MAX
    Xs = Zfull[:, keep]

    rng = np.random.RandomState(RANDOM_SEED)
    hdr = ('target n |tau| cb 95% CI |tau| raw |tau| loc frac>0 agr')
    print(hdr)

    gates = []
    for name, (Xt, yt, sidt, agt, ivt) in targets.items():
        n_clips, n_spk = len(yt), int(np.unique(sidt).size)
        Xc = Xt[:, keep]

        if fig == "a":
            r = _ratio(ags, agt)
        elif fig == "b":
            r = np.ones(len(ys))
        else:
            Z_src = np.where(np.isnan(drs_nv), 0.0, drs_nv)
            Z_t = ivt
            r = np.ones(len(ys))
            for c in np.unique(ys):
                ms = ys == c
                mt = (yt == c)
                if ms.sum() > 1 and mt.sum() > 1 and np.std(Z_t[mt]) > 1e-9:
                    r[ms] = _ratio(Z_src[ms], Z_t[mt])
        if np.any(np.isnan(r)):
            r = np.nan_to_num(r, nan=1.0)

        Z_src = ags if fig == "a" else (None if fig == "b" else
                                        np.where(np.isnan(drs_nv), 0.0, drs_nv))
        B = _run_target(Xs, ys, ags, pids, Xc, yt, sidt, agt, r, rng, N_BOOT,
                        model=cfg["graph"], ups=None, grs=None, z=Z_src)
        cb, raw, loc = B["cb"], B["raw"], B["loc"]
        rm = lambda a: np.sqrt(np.mean(a ** 2, axis=1))
        crm, rrm, lrm = rm(cb), rm(raw), rm(loc)
        lo, hi = np.percentile(crm, 2.5), np.percentile(crm, 97.5)

        W_all = _cb(Xs, ys, ags, model=cfg["graph"],
                    z=(ags if fig == "a" else (None if fig == "b" else
                                               np.where(np.isnan(drs_nv), 0.0, drs_nv))))
        cb_pt = _effect_vec(Xs, W_all[:, 1] * r, W_all[:, 0] * r)
        loc_pt = _local_effect(Xc, yt, agt)
        agr = _sign_agr(cb_pt, loc_pt)

        lo_f = np.percentile(cb, 2.5, axis=0)
        hi_f = np.percentile(cb, 97.5, axis=0)
        frac0 = float(np.mean((lo_f > 0) | (hi_f < 0)))

        dd = crm - rrm
        pval = 1.0
        if np.any(dd) and len(np.unique(dd)) > 1:
            try:
                _, pval = wilcoxon(dd)
            except ValueError:
                pval = 1.0

        print(str(name) + ' ' + str(n_clips) + ' ' + str(np.median(crm)) + '[' + str(lo) + ',' + str(hi) + ']' + str(np.median(rrm)) + ' ' + str(np.median(lrm)) + ' ' + str(frac0) + ' ' + str(agr), flush=True)
        gates.append(dict(target=name, lo=lo, hi=hi, agr=agr,
                          p_deconf=pval))

    print()
    print("Gates (fig " + fig + ", pre-committed in tests/gates.txt):")
    n1 = sum(1 for g in gates if g["lo"] > 0)
    n2 = sum(1 for g in gates if g["agr"] >= 0.5)
    n3 = sum(1 for g in gates if g["p_deconf"] < 0.05)
    print(' g1 |tau|_cb 95%-ci lower > 0 on ' + str(n1) + '/' + str(len(gates)) + ' targets')
    print(' g2 sign agreement cb vs localized >= .5 on ' + str(n2) + '/' + str(len(gates)) + ' targets')
    print(' g3 cb differs from raw assoc (p<.05) on ' + str(n3) + '/' + str(len(gates)) + ' targets')

if __name__ == "__main__":
    main()
