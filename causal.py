import sys, os, re
import numpy as np
import pandas as pd
from pathlib import Path

BASE = Path(__file__).resolve().parent
for p in [str(BASE) + '/libraries/CausalBootstrapping', str(BASE) + '/libraries/GRAPL',
          str(BASE) + '/libraries/causal_assistant', str(BASE) + '/tCB', str(BASE) + '/tests']:
    if p not in sys.path:
        sys.path.insert(0, p)

import causal_assistant as ca
from scipy.stats import gaussian_kde
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler as _SS

import transport_u_age as age

RANDOM_SEED = 42
N_REPS = 200
U_DIMS = 2
N_AGE_BINS = 6
CB_FIT_METHOD = "kde"
META_COLS = {"Unnamed: 0", "tot_name"}
NV_FEAT = str(BASE) + '/data/neurovoz/extracted_features'
INTENSITY_IDX = slice(0, 38)

AGE_GRAPH = "y;X;Age; y->X; Age->X; Age->y;"
UPDRS_EDGES = [0, 10, 20, 40, 70, 200]
GRBAS_EDGES = [0, 1, 2, 3, 4]
GRAPH_MODELS = {
    "age":   "y;X;Age; y->X; Age->X; Age->y;",
    "grbas": "y;X;Age;GRBAS; y->X; Age->X; Age->y; y->GRBAS; GRBAS->X;",
    "updrs": "y;X;Age;UPDRS; y->X; Age->X; Age->y; y->UPDRS; UPDRS->X;",
    "both":  "y;X;Age;UPDRS;GRBAS; y->X; Age->X; Age->y; y->UPDRS; UPDRS->X; "
             "y->GRBAS; GRBAS->X;",
    "g-as[a]":   "y;X;Age;Sex;GRBAS; y->X; Age->X; Age->y; Sex->X; Sex->y; "
                 "y->GRBAS; GRBAS->X;",
    "ug-as[a]":  "y;X;Age;Sex;UPDRS;GRBAS; y->X; Age->X; Age->y; Sex->X; Sex->y; "
                 "y->UPDRS; UPDRS->X; y->GRBAS; GRBAS->X;",
    "xy":    "y;X; X->y;",
}
INTERVALS = {"-2": (-10, -6), "-1": (-5, -1), "1": (0, 5), "2": (6, 10)}

SYNTH_QY_CONF = 0.80
SYNTH_QY_UNCONF = 0.50
SYNTH_QY_REV = 0.20
SYNTH_QYZ = 0.95
SYNTH_BRIGHT = 1.0

def _balanced_acc(y, pred):
    p0 = np.mean(pred[y == 0] == 0) if np.any(y == 0) else np.nan
    p1 = np.mean(pred[y == 1] == 1) if np.any(y == 1) else np.nan
    return 0.5 * (p0 + p1)

def _f1_score(y, pred):
    tp = np.sum((pred == 1) & (y == 1))
    fp = np.sum((pred == 1) & (y == 0))
    fn = np.sum((pred == 0) & (y == 1))
    if tp == 0:
        return 0.0
    p, r = tp / (tp + fp), tp / (tp + fn)
    return 2 * p * r / (p + r) if p + r > 0 else 0.0

def _roc_auc(y, prob):
    y = np.asarray(y)
    pos = y == 1
    n_pos, n_neg = int(pos.sum()), int((~pos).sum())
    if n_pos == 0 or n_neg == 0:
        return np.nan
    order = np.argsort(prob, kind="mergesort")
    ranks = np.empty(len(prob), dtype=float)
    ranks[order] = np.arange(1, len(prob) + 1)
    return (ranks[pos].sum() - n_pos * (n_pos + 1) / 2) / (n_pos * n_neg)

def _metrics(y, prob):
    y = np.asarray(y)
    yi = y.astype(int)
    pred = (np.asarray(prob) >= 0.5).astype(int)
    acc = np.mean(pred == yi)
    ba = _balanced_acc(yi, pred)
    f1v = _f1_score(yi, pred)
    auc = _roc_auc(yi, prob)
    return acc, ba, f1v, auc

def _nan_mean(a, axis=0):
    return np.nanmean(np.asarray(a, dtype=float), axis=axis)

def _nan_std(a, axis=0):
    return np.nanstd(np.asarray(a, dtype=float), axis=axis)

def _stem(x):
    s = str(x).rsplit("/", 1)[-1]
    while s.endswith(".wav"):
        s = s[:-4]
    return s.replace("audios-", "").replace("audios/", "")

def _sl(df, kv, common, meta_cols):
    return (df.loc[kv.isin(common)]
              .loc[:, ~df.columns.isin(meta_cols)]
              .set_index(kv[kv.isin(common)].values))

def _task(c):
    s = str(c)
    s = s.rsplit("/", 1)[-1]
    while s.endswith(".wav"):
        s = s[:-4]
    s = s.replace("audios-", "").replace("audios/", "")
    if s.startswith("PD_"):
        s = s[3:]
    elif s.startswith("HC_"):
        s = s[3:]
    s = re.sub(r"_\d+$", "", s)
    return s

def load_nv(task=None):
    hc = pd.read_csv(str(BASE) + '/data/neurovoz/metadata/metadata_hc.csv')
    pd_ = pd.read_csv(str(BASE) + '/data/neurovoz/metadata/metadata_pd.csv')
    mt = pd.concat([hc, pd_], ignore_index=True)
    mt["_pid"] = mt["ID"].map(lambda i: str(int(i)).zfill(4))
    mt["Age"] = pd.to_numeric(mt["Age"], errors="coerce")
    a_by_pid = mt.groupby("_pid")["Age"].median()
    mt["Sex"] = pd.to_numeric(mt["Sex"], errors="coerce")
    s_by_pid = mt.groupby("_pid")["Sex"].median()
    mt["UPDRS scale"] = pd.to_numeric(mt["UPDRS scale"], errors="coerce")
    u_by_pid = mt.groupby("_pid")["UPDRS scale"].median()

    grbas = {}
    for f in os.listdir(str(BASE) + '/data/grbas'):
        if not f.endswith(".csv"):
            continue
        for _, r in pd.read_csv(str(BASE) + '/data/grbas/' + f).iterrows():
            gid = os.path.splitext(os.path.basename(str(r["text_patient_id"])))[0].strip()
            try:
                grbas[gid] = float(r["G"])
            except (ValueError, TypeError):
                pass

    intens = pd.read_csv(str(NV_FEAT) + '/intensity_features_neurovoz.csv')
    pause = pd.read_csv(str(NV_FEAT) + '/pause_features_neurovoz.csv')
    syll = pd.read_csv(str(NV_FEAT) + '/syllable_features_neurovoz.csv')
    ki, kp, ks = (intens["tot_name"].apply(_stem),
                  pause["tot_name"].apply(_stem),
                  syll["tot_name"].apply(_stem))
    common = sorted(set(ki) & set(kp) & set(ks))
    if task is not None:
        common = [c for c in common if _task(c) == task]

    X = pd.concat([_sl(intens, ki, common, META_COLS),
                   _sl(pause, kp, common, META_COLS),
                   _sl(syll, ks, common, META_COLS)], axis=1)
    X = X.values.astype(np.float64)
    Y = np.array([1 if c.startswith("PD_") else 0 for c in common], dtype=np.float64)
    pid = np.array([c[-4:] for c in common])
    ag = np.array([a_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    sx = np.array([s_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    up = np.array([np.nan if np.isnan(u_by_pid.get(p, np.nan)) else
                   u_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    gr = np.array([grbas.get(c[3:], np.nan) for c in common], dtype=np.float64)
    ok = ~(np.isnan(X).any(axis=1) | np.isnan(ag) | np.isnan(sx))
    X, Y, ag, sx, pid, up, gr = X[ok], Y[ok], ag[ok], sx[ok], pid[ok], up[ok], gr[ok]
    if task is not None:
        return X, Y, ag, sx, pid, up, gr
    tasks = np.array([_task(c) for c in common])[ok]
    return X, Y, ag, sx, pid, up, gr, tasks

def load_pc():
    base = str(BASE) + '/data/parkceleb_extracted'
    intens = pd.read_csv(str(base) + '/intensity_features_parkceleb.csv')
    pause = pd.read_csv(str(base) + '/pause_features_parkceleb.csv')
    syll = pd.read_csv(str(base) + '/syllable_features_parkceleb.csv')
    kid = intens["audio_file"].map(
        lambda f: (lambda m: m.group(1) if m else np.nan)(
            age.IDRE.search(str(f))))
    m = kid.notna()
    intens, pause, syll, kid = (intens[m].reset_index(drop=True),
                                pause[m].reset_index(drop=True),
                                syll[m].reset_index(drop=True),
                                kid[m].reset_index(drop=True))
    X = pd.concat([intens.loc[:, ~intens.columns.isin(age.FEATURE_META)],
                   pause.loc[:, ~pause.columns.isin(age.FEATURE_META)],
                   syll.loc[:, ~syll.columns.isin(age.FEATURE_META)]], axis=1)
    X = X.values.astype(np.float64)
    Y = intens["parkceleb_label"].values.astype(np.float64)
    sid = intens["speaker_id"].astype(str).values

    ages = pd.read_csv(str(BASE) + '/data/parkceleb_ages.csv')
    age_by_af = (ages.drop_duplicates("audio_file").set_index("audio_file")["age"])
    af = intens["audio_file"].astype(str).map(
        lambda f: f if f in age_by_af.index else np.nan)
    t_age = af.map(age_by_af).values.astype(np.float64)

    upload = pd.read_csv(str(BASE) + '/data/parkceleb_upload_years.csv')
    upload["video_id"] = upload["video_id"].astype(str)
    up_full = upload["upload_date"].astype(str).str[:4]
    up = up_full[up_full.str.fullmatch(r"\d{4}")].astype(int)
    up.index = upload.loc[up.index, "video_id"]

    import openpyxl
    wb = openpyxl.load_workbook(str(BASE) + '/data/ParkCeleb/PD_demo.xlsx',
                                read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb2 = openpyxl.load_workbook(str(BASE) + '/data/ParkCeleb/speakers_pairs.xlsx',
                                 read_only=True)
    rows2 = list(wb2.active.iter_rows(values_only=True))
    pairs = {}
    for r in rows2[1:]:
        if r[0] is None:
            continue
        pairs[str(r[0]).strip()] = str(r[1]).strip()

    iv = np.full(len(X), np.nan)
    for i, sp in enumerate(sid):
        vid = kid[i]
        if vid not in up.index:
            continue
        uy = up[vid]
        yr = None
        iv[i] = uy - yr if yr is not None else np.nan

    out = {"full": (X, Y, sid, t_age)}
    for name, (lo, hi) in INTERVALS.items():
        sel = (iv >= lo) & (iv <= hi)
        out[name] = (X[sel], Y[sel], sid[sel], t_age[sel])
    return out

def _gbc(X, y, w=None):
    clf = GradientBoostingClassifier(n_estimators=100, learning_rate=0.1,
                                     max_depth=3, random_state=RANDOM_SEED)
    if w is not None:
        w = np.asarray(w, dtype=float)
        w = w * (len(X) / w.sum())
    clf.fit(np.asarray(X), y.astype(int), sample_weight=w)
    return clf

def _gbc_two_stage(X_base, y_base, w_base, X_spec, y_spec, w_spec, extra_trees=20):
    clf = GradientBoostingClassifier(n_estimators=100, learning_rate=0.1,
                                     max_depth=3, random_state=RANDOM_SEED,
                                     warm_start=True)
    if w_base is not None:
        w_base = np.asarray(w_base, dtype=float)
        w_base = w_base * (len(X_base) / w_base.sum())
    clf.fit(np.asarray(X_base), y_base.astype(int), sample_weight=w_base)
    
    clf.n_estimators += extra_trees
    if w_spec is not None:
        w_spec = np.asarray(w_spec, dtype=float)
        w_spec = w_spec * (len(X_spec) / w_spec.sum())
    clf.fit(np.asarray(X_spec), y_spec.astype(int), sample_weight=w_spec)
    
    return clf

def _cb(X, y, ag, sx, model="age", up=None, gr=None):
    graph = GRAPH_MODELS[model]
    if "X->y" in graph and "y->X" not in graph:
        return np.ones((len(X), 2))
    fit = CB_FIT_METHOD
    kwargs = dict(causal_graph=graph, cause_var="y",
                  effect_var="X", fit_method=fit,
                  X=X, y=y, Age=(ag, N_AGE_BINS))
    if "Sex" in graph:
        if fit == "kde":
            jit = np.random.default_rng(len(X)).uniform(-0.005, 0.005, len(sx))
            sx = np.asarray(sx, dtype=float) + jit
        kwargs["Sex"] = (sx, 2)
    if "UPDRS" in graph:
        u = np.asarray(up, dtype=float)
        pd_mask = y == 1
        u_pd_med = np.nanmedian(u[pd_mask]) if np.any(np.isfinite(u[pd_mask])) else 0.0
        u = np.where(np.isnan(u), np.where(pd_mask, u_pd_med, 0.0), u)
        kwargs["UPDRS"] = (u, len(UPDRS_EDGES) - 1)
    if "GRBAS" in graph:
        g = np.asarray(gr, dtype=float)
        pd_mask = y == 1
        g_pd_med = np.nanmedian(g[pd_mask]) if np.any(np.isfinite(g[pd_mask])) else 0.0
        g = np.where(np.isnan(g), np.where(pd_mask, g_pd_med, 0.0), g)
        kwargs["GRBAS"] = (g, len(GRBAS_EDGES) - 1)
    return np.asarray(ca.causal_weights(**kwargs))

def _tcb_weights(U_src, U_trg, Y_trg):
    U_src = np.asarray(U_src, dtype=float)
    U_trg = np.asarray(U_trg, dtype=float)
    if U_src.ndim == 1:
        U_src = U_src[:, None]
    if U_trg.ndim == 1:
        U_trg = U_trg[:, None]
    pu_src = np.asarray(gaussian_kde(U_src.T)(U_trg.T)).ravel()
    pu_trg = np.asarray(gaussian_kde(U_trg.T)(U_trg.T)).ravel()
    lr = LogisticRegression(C=1e6, solver="lbfgs", max_iter=5000, random_state=42)
    lr.fit(U_trg, Y_trg.ravel())
    py1 = lr.predict_proba(U_trg)[:, 1]
    py_given_u = np.where(Y_trg.ravel() > 0.5, py1, 1 - py1)
    w = pu_src / (pu_trg * py_given_u)
    return w / np.mean(w)

def _loc_weights(X, y, age):
    ag = np.asarray(age, dtype=float)
    med = np.nanmedian(ag) if np.any(np.isfinite(ag)) else 0.0
    ag = np.where(np.isfinite(ag), ag, med)
    W = np.asarray(ca.causal_weights(causal_graph=AGE_GRAPH, cause_var="y",
                                     effect_var="X", fit_method=CB_FIT_METHOD,
                                     X=X, y=y, Age=(ag, N_AGE_BINS)))
    return W

def _w_backdoor(u, y):
    u = np.asarray(u, dtype=float)
    k_all = gaussian_kde(u)
    p_all = k_all(u)
    w = np.empty(len(u))
    for c in (0, 1):
        m = y == c
        k_c = gaussian_kde(u[m])
        w[m] = p_all[m] / np.maximum(k_c(u[m]), 1e-12)
    return w * (len(u) / w.sum())

def _synth_conf(X_pool, y_pool, quy, qyz, bright, seed_rand, seed_randn, frac=0.95):
    rng_rand = np.random.default_rng(seed_rand)
    rng_randn = np.random.default_rng(seed_randn)
    N = X_pool.shape[0]
    perm = rng_rand.permutation(N)
    Xp, yp = X_pool[perm], y_pool[perm]
    M = int(round(N * frac))
    fstd = np.maximum(np.std(Xp, 0), 1e-12)

    U = rng_randn.normal(0.0, 5.0, M)
    q = np.log((1 - quy) / quy)
    p_y = 1 / (1 + np.exp(q * U))
    Y = (rng_rand.uniform(size=M) <= p_y).astype(int)

    Z = np.where(rng_rand.uniform(size=M) > qyz, 1 - Y, Y)

    idx = {c: np.where(yp == c)[0] for c in (0, 1)}
    j = np.ones(2, dtype=int)
    X = np.empty((M, Xp.shape[1]), dtype=np.float64)
    for n in range(M):
        g = Z[n]
        cur = (j[g] - 1) % len(idx[g])
        i = idx[g][cur]
        j[g] += 1
        X[n] = Xp[i] + (0.5 * np.tanh(0.2 * U[n]) + 0.5) * bright * fstd
    return X, Y, Z, U

def _speaker_split(sid, y, seed):
    rng = np.random.default_rng(seed)
    uniq = np.unique(sid)
    maj = np.array([np.bincount(y[sid == u].astype(int)).argmax() for u in uniq])
    tr_u, te_u = [], []
    for c in (0, 1):
        idx = np.where(maj == c)[0]
        p = rng.permutation(idx)
        k = int(round(0.8 * len(p)))
        tr_u += list(uniq[p[:k]].tolist())
        te_u += list(uniq[p[k:]].tolist())
    tr = np.isin(sid, np.asarray(tr_u))
    te = np.isin(sid, np.asarray(te_u))
    return tr, te

def _perm_gt(a, b, n_flip=10000, seed=7):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    d = a - b
    d = d[np.isfinite(d)]
    if len(d) < 3 or np.std(d) == 0:
        return 1.0
    rng = np.random.default_rng(seed)
    obs = d.mean()
    flips = rng.integers(0, 2, size=(n_flip, len(d)), dtype=np.int8)
    null = (d[None, :] * (2 * flips - 1)).mean(axis=1)
    return float((null >= obs).mean())

def _backdoor_effect(U_tr, v_tr, y_tr, U_tgt, reps=N_REPS):
    rng = np.random.default_rng(RANDOM_SEED)
    rows = [[], [], []]
    for rep in range(reps):
        n = len(v_tr)
        idx = rng.integers(0, n, n)
        u, v, y = U_tr[idx], v_tr[idx], y_tr[idx]
        k_all = gaussian_kde(u.T)(u.T)
        k_y = np.empty(n)
        for c in (0, 1):
            m = y == c
            if m.sum() < 3:
                return rows
            k_y[m] = gaussian_kde(u[m].T)(u[m].T)
        w_s = k_all / np.maximum(k_y, 1e-12)
        k_t = gaussian_kde(U_tgt[rng.integers(0, len(U_tgt), n)].T)(u.T)
        w_t = k_t / np.maximum(k_y, 1e-12)
        rows[0].append(v[y == 1].mean() - v[y == 0].mean())
        rows[1].append(w_s[y == 1].dot(v[y == 1]) / w_s[y == 1].sum()
                       - w_s[y == 0].dot(v[y == 0]) / w_s[y == 0].sum())
        rows[2].append(w_t[y == 1].dot(v[y == 1]) / w_t[y == 1].sum()
                       - w_t[y == 0].dot(v[y == 0]) / w_t[y == 0].sum())
    return rows

def _eval_backdoor(U_nv, v_nv, ys, X_pc, mu_s, sd_s, pca, k_u):
    m = ~np.isnan(X_pc).any(axis=1)
    U_pc = pca.transform((X_pc[m][:, INTENSITY_IDX] - mu_s[INTENSITY_IDX])
                         / sd_s[INTENSITY_IDX])[:, k_u]
    rows = _backdoor_effect(U_nv, v_nv, ys, U_pc)
    print()
    print("step 3: causal effect of pd on voice, backdoor over u")
    print("   u -> y, u -> x, y -> x   (u measured on both corpora)")
    hdr = "estimator          effect              95% ci    population"
    print(hdr)
    for lab, k in (("naive", 0), ("backdoor, source p(u)", 1),
                   ("backdoor, target p(u)", 2)):
        a = np.asarray(rows[k], dtype=float)
        a = a[np.isfinite(a)]
        lo, hi = np.percentile(a, 2.5), np.percentile(a, 97.5)
        pop = "pd under nv severity" if k < 2 else "pd under pc severity"
        print(lab + "  " + str(round(a.mean(), 4)) + "  [" + str(round(lo, 4)) + ", " + str(round(hi, 4)) + "]  " + pop)
    return rows

def _eval_pc(X_pc, y_pc, sid_pc, age_pc, U_nv, pca, k_u, label):
    m0 = ~np.isnan(X_pc).any(axis=1)
    X_pc, y_pc, sid_pc, age_pc = X_pc[m0], y_pc[m0], sid_pc[m0], age_pc[m0]
    if len(np.unique(sid_pc)) < 4:
        print()
        print(+ label + ": only " + str(len(np.unique(sid_pc))) + " speakers, skipped")
        return
    arms = ["standard", "localized", "tcb"]
    res = {k: [] for k in arms}
    for rep in range(N_REPS):
        seed = RANDOM_SEED + rep
        tr, te = _speaker_split(sid_pc, y_pc, seed)
        if tr.sum() == 0 or te.sum() == 0:
            continue
        sc_pc = _SS().fit(X_pc[tr])
        X_tr, X_te = sc_pc.transform(X_pc[tr]), sc_pc.transform(X_pc[te])
        y_tr, y_te = y_pc[tr], y_pc[te]
        U_pc_tr = pca.transform(X_tr[:, INTENSITY_IDX])[:, k_u]
        w_loc = _loc_weights(X_tr, y_tr, age_pc[tr])
        w_loc = w_loc[np.arange(len(y_tr)), y_tr.astype(int)]
        w_tcb = _tcb_weights(U_nv, U_pc_tr, y_tr)
        g_std = _gbc(X_tr, y_tr)
        g_loc = _gbc(X_tr, y_tr, w_loc)
        g_tcb = _gbc(X_tr, y_tr, w_tcb)
        res["standard"].append(_metrics(y_te, g_std.predict_proba(X_te)[:, 1]))
        res["localized"].append(_metrics(y_te, g_loc.predict_proba(X_te)[:, 1]))
        res["tcb"].append(_metrics(y_te, g_tcb.predict_proba(X_te)[:, 1]))

    print()
    print(" + label + ": n=" + str(len(y_pc)) + ", speakers=" + str(len(np.unique(sid_pc))) + ", reps=" + str(N_REPS) + ")
    hdr = "method                     acc                  ba                   f1                 auc     p_pair               diff_ci"
    print(hdr)
    for key, lab in (("standard", "standard"),
                     ("localized", "localised (pc)"),
                     ("tcb", "tcb (nv->pc)")):
        a = np.asarray(res[key])
        row = ""
        for k in range(4):
            row += str(round(_nan_mean(a[:, k]), 4)) + "+/-" + str(round(_nan_std(a[:, k]), 4)) + "  "
        if key == "standard":
            print(lab + "  " + row)
        else:
            ref = np.asarray(res["standard"])
            d = a[:, 1] - ref[:, 1]
            d = d[np.isfinite(d)]
            lo, hi = np.percentile(d, 2.5), np.percentile(d, 97.5)
            p_pair = _perm_gt(a[:, 1], ref[:, 1])
            print(lab + "  " + row + "  " + str(round(p_pair, 4)) + "  [" + str(round(lo, 4)) + ", " + str(round(hi, 4)) + "]")

def main():
    X_raw, ys, ags, sxs, _, ups, grs, tasks = load_nv()
    X_raw_pb, ys_pb, ags_pb, sxs_pb, _, ups_pb, grs_pb = load_nv(task="PETACA_BLANCA")
    if ags.size > 0:
        age_str = "age [" + str(int(np.nanmin(ags))) + "-" + str(int(np.nanmax(ags))) + "]"
    else:
        age_str = "age [no data]"
    print("nv: " + str(X_raw.shape[0]) + " recordings (" + str(int(ys.sum())) + " pd / " + str(len(ys) - int(ys.sum())) + " hc), " + age_str)

    mu_s, sd_s = X_raw.mean(0), X_raw.std(0)
    sd_s = np.maximum(sd_s, np.median(sd_s) * 0.05)
    Xs = (X_raw - mu_s) / sd_s
    Xs_pb = (X_raw_pb - mu_s) / sd_s

    print()
    print("step 1: semi-synthetic confounder verification (simulated u)")
    sc = _SS().fit(X_raw)
    Xc = sc.transform(X_raw)
    tr_conf = _synth_conf(Xc, ys, SYNTH_QY_CONF, SYNTH_QYZ, SYNTH_BRIGHT, 1, 101)
    tr_unconf = _synth_conf(Xc, ys, SYNTH_QY_UNCONF, SYNTH_QYZ, SYNTH_BRIGHT, 2, 202)
    tr_rev = _synth_conf(Xc, ys, SYNTH_QY_REV, SYNTH_QYZ, SYNTH_BRIGHT, 5, 505)
    tr_proj = _synth_conf(Xc, ys, SYNTH_QY_CONF, SYNTH_QYZ, 0.0, 6, 606)
    te_conf = _synth_conf(Xc, ys, SYNTH_QY_CONF, SYNTH_QYZ, SYNTH_BRIGHT, 3, 303)
    te_unconf = _synth_conf(Xc, ys, SYNTH_QY_UNCONF, SYNTH_QYZ, SYNTH_BRIGHT, 4, 404)

    w_dec = _w_backdoor(tr_conf[3], tr_conf[1])
    w_dec_r = _w_backdoor(tr_rev[3], tr_rev[1])
    w_proj = _w_backdoor(tr_proj[3], tr_proj[1])
    mods = {
        "confounded": _gbc(tr_conf[0], tr_conf[1]),
        "deconfounded": _gbc(tr_conf[0], tr_conf[1], w_dec),
        "unconf-trained": _gbc(tr_unconf[0], tr_unconf[1]),
        "confounded (reverse)": _gbc(tr_rev[0], tr_rev[1]),
        "deconfounded (reverse)": _gbc(tr_rev[0], tr_rev[1], w_dec_r),
        "deconfounded (u-proj)": _gbc(tr_proj[0], tr_proj[1], w_proj),
    }
    print("                 model     conf-test acc     conf-test ba    unconf-test acc   unconf-test ba")
    for name, g in mods.items():
        m1 = _metrics(te_conf[1], g.predict_proba(te_conf[0])[:, 1])
        m2 = _metrics(te_unconf[1], g.predict_proba(te_unconf[0])[:, 1])
        print(name.rjust(22) + "  " + str(round(m1[0], 4)) + "  " + str(round(m1[1], 4)) + "  " + str(round(m2[0], 4)) + "  " + str(round(m2[1], 4)))

    pc = load_pc()
    print()
    print("step 2: causal transportability nv -> pc (speaker-level split)")

    pca = PCA(n_components=5)
    L_nv = pca.fit_transform(Xs[:, INTENSITY_IDX])
    corrs_u = np.array([abs(np.corrcoef(L_nv[:, i], ys)[0, 1]) for i in range(5)])
    k_u = np.argsort(corrs_u)[:U_DIMS]
    U_nv = L_nv[:, k_u]
    print("   u = pcs " + str(k_u + 1) + " of nv intensity block (|r(y)| <= " + str(round(corrs_u[k_u].max(), 3)) + ")")

    for name, (X_pc, y_pc, sid_pc, age_pc) in pc.items():
        _eval_pc(X_pc, y_pc, sid_pc, age_pc, U_nv, pca, k_u, name)

    v_pc = PCA(n_components=1).fit(Xs).transform(Xs)[:, 0]
    _eval_backdoor(U_nv, v_pc, ys, pc["full"][0], mu_s, sd_s, pca, k_u)

    print()
    print("step 2b: two-stage transfer learning (all tasks -> petaca blanca)")
    print("base source: all tasks (n=" + str(len(ys)) + ")")
    print("fine-tune source: petaca blanca (n=" + str(len(ys_pb)) + ")")

    for model in GRAPH_MODELS:
        W_base = _cb(Xs, ys, ags, sxs, model=model, up=ups, gr=grs)
        w_base = W_base[np.arange(len(ys)), ys.astype(int)]
        
        W_spec = _cb(Xs_pb, ys_pb, ags_pb, sxs_pb, model=model, up=ups_pb, gr=grs_pb)
        w_spec = W_spec[np.arange(len(ys_pb)), ys_pb.astype(int)]
        
        g_cb = _gbc_two_stage(Xs, ys, w_base, Xs_pb, ys_pb, w_spec, extra_trees=20)
        g_raw = _gbc_two_stage(Xs, ys, None, Xs_pb, ys_pb, None, extra_trees=20)
        
        for name, (X_pc, y_pc, _, _) in pc.items():
            m0 = ~np.isnan(X_pc).any(axis=1)
            Xt = (X_pc[m0] - mu_s) / sd_s
            yt = y_pc[m0]
            mc = _metrics(yt, g_cb.predict_proba(Xt)[:, 1])
            mr = _metrics(yt, g_raw.predict_proba(Xt)[:, 1])
            print(model.rjust(9) + " " + name.rjust(6) + "  " + str(round(mc[0], 4)) + "/" + str(round(mr[0], 4)) + " " + str(round(mc[1], 4)) + "/" + str(round(mr[1], 4)) + " " + str(round(mc[2], 4)) + "/" + str(round(mr[2], 4)) + " " + str(round(mc[3], 4)) + "/" + str(round(mr[3], 4)))

if __name__ == "__main__":
    main()