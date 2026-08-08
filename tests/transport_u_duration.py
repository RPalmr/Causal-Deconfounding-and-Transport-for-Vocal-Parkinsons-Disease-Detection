import sys, os
import numpy as np
import pandas as pd
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for p in [f"{BASE}/libraries/CausalBootstrapping", f"{BASE}/libraries/GRAPL",
          f"{BASE}/libraries/causal_assistant", f"{BASE}/tCB", f"{BASE}/tests"]:
    if p not in sys.path:
        sys.path.insert(0, p)

import causal_assistant as ca

from sklearn.ensemble import GradientBoostingClassifier, ExtraTreesClassifier
from sklearn.metrics import roc_auc_score, f1_score

import transport_u_age as age
import transport_u_duration as dur

RANDOM_SEED = 42
N_BOOT = 200
AGE_EDGES = np.arange(40, 101, 10)
META_COLS = {"Unnamed: 0", "tot_name"}
INTERVALS = {"-2": (-10, -6), "-1": (-5, -1), "1": (0, 5), "2": (6, 10)}
NV_FEAT = f"{BASE}/data/neurovoz/extracted_features"
N_SELECT = 30
UPDRS_EDGES = [0, 10, 20, 40, 70, 200]
GRBAS_EDGES = [0, 1, 2, 3, 4]
GBM_GRID = {"n_estimators": [100, 200], "max_depth": [3, 5],
            "learning_rate": [0.05, 0.1]}
CB_FIT_METHOD = "kde"

UGAS_GRAPH = "y;X;Age;Sex;UPDRS;GRBAS; y<->X; y->UPDRS; UPDRS->X; y->GRBAS; GRBAS->X; Age->X; Age->y; Sex->X; Sex->y;"

def load_nv_petaca():
    hc = pd.read_csv(f"{BASE}/data/neurovoz/metadata/metadata_hc.csv")
    pd_ = pd.read_csv(f"{BASE}/data/neurovoz/metadata/metadata_pd.csv")
    mt = pd.concat([hc, pd_], ignore_index=True)
    mt["_pid"] = mt["ID"].map(lambda i: f"{int(i):04d}")
    mt["Age"] = pd.to_numeric(mt["Age"], errors="coerce")
    mt["Sex"] = pd.to_numeric(mt["Sex"], errors="coerce")
    mt["UPDRS scale"] = pd.to_numeric(mt["UPDRS scale"], errors="coerce")
    a_by_pid = mt.groupby("_pid")["Age"].median()
    s_by_pid = mt.groupby("_pid")["Sex"].median()
    u_by_pid = mt.groupby("_pid")["UPDRS scale"].median()

    grbas = {}
    for f in os.listdir(f"{BASE}/data/grbas"):
        if not f.endswith(".csv"):
            continue
        for _, r in pd.read_csv(f"{BASE}/data/grbas/{f}").iterrows():
            gid = os.path.splitext(os.path.basename(str(r["text_patient_id"])))[0].strip()
            try:
                grbas[gid] = float(r["G"])
            except (ValueError, TypeError):
                pass

    intens = pd.read_csv(f"{NV_FEAT}/intensity_features_neurovoz.csv")
    pause = pd.read_csv(f"{NV_FEAT}/pause_features_neurovoz.csv")
    syll = pd.read_csv(f"{NV_FEAT}/syllable_features_neurovoz.csv")
    def stem(x):
        s = os.path.basename(str(x))
        while s.endswith(".wav"):
            s = s[:-4]
        return s.replace("audios-", "").replace("audios/", "")
    ki, kp, ks = (intens["tot_name"].apply(stem),
                  pause["tot_name"].apply(stem),
                  syll["tot_name"].apply(stem))
    common = sorted(set(ki) & set(kp) & set(ks))
    common = [c for c in common if "PETACA" in c]
    X = pd.concat([dur._sl(intens, ki, common, META_COLS),
                   dur._sl(pause, kp, common, META_COLS),
                   dur._sl(syll, ks, common, META_COLS)], axis=1)
    X = X.values.astype(np.float64)
    Y = np.array([1 if c.startswith("PD_") else 0 for c in common], dtype=np.float64)
    pid = np.array([c[-4:] for c in common])
    ag = np.array([a_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    sx = np.array([s_by_pid.get(p, np.nan) for p in pid], dtype=np.float64)
    up = np.array([u_by_pid.get(p, 0.0) if not np.isnan(u_by_pid.get(p, np.nan)) else 0.0
                   for p in pid], dtype=np.float64)
    up = np.where(np.isnan(up), 0.0, up)
    gr = np.array([grbas.get(c[3:], np.nan) for c in common], dtype=np.float64)
    ok = ~(np.isnan(X).any(axis=1) | np.isnan(ag) | np.isnan(sx) | np.isnan(gr))
    return X[ok], Y[ok], ag[ok], sx[ok], up[ok], gr[ok], pid[ok]

def load_pc_intervals():
    base = f"{BASE}/data/parkceleb_extracted"
    intens = pd.read_csv(f"{base}/intensity_features_parkceleb.csv")
    pause = pd.read_csv(f"{base}/pause_features_parkceleb.csv")
    syll = pd.read_csv(f"{base}/syllable_features_parkceleb.csv")
    kid = intens["audio_file"].map(
        lambda f: (lambda m: m.group(1) if m else np.nan)(
            age.IDRE.search(str(f))))
    m = kid.notna()
    intens, pause, syll, kid = (intens[m].reset_index(drop=True),
                                pause[m].reset_index(drop=True),
                                syll[m].reset_index(drop=True),
                                kid[m].reset_index(drop=True))
    X = pd.concat([intens.loc[:, ~intens.columns.isin(age.FEATURE_META)],
                   pause.loc[:, ~pause.columns.isin(age.FEATURE_META)],
                   syll.loc[:, ~syll.columns.isin(age.FEATURE_META)]], axis=1)
    X = X.values.astype(np.float64)
    Y = intens["parkceleb_label"].values.astype(np.float64)
    sid = intens["speaker_id"].astype(str).values

    ages_csv = pd.read_csv(f"{BASE}/data/parkceleb_ages.csv")
    age_by_af = (ages_csv.drop_duplicates("audio_file")
                 .set_index("audio_file")["age"])
    af_key = intens["audio_file"].astype(str).map(
        lambda f: f if f in age_by_af.index else np.nan)
    t_age_all = af_key.map(age_by_af).values.astype(np.float64)

    def _demo_gender(path, ncols):
        wb = openpyxl.load_workbook(f"{BASE}/{path}", read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        hdr = [str(c).strip().lower().replace(" ", "") for c in rows[0]]
        gi = hdr.index("gender")
        out = {}
        for r in rows[2:]:
            if r[0] is None:
                continue
            out[str(r[0]).strip()] = (1.0 if str(r[gi]).strip().lower()
                                      .startswith("m") else 0.0)
        return out
    gender = {}
    gender.update(_demo_gender("data/ParkCeleb/PD_demo.xlsx", 5))
    gender.update(_demo_gender("data/ParkCeleb/CN_demo.xlsx", 4))
    t_sex_all = np.array([gender.get(s, np.nan) for s in sid],
                         dtype=np.float64)

    upload = pd.read_csv(f"{BASE}/data/parkceleb_upload_years.csv")
    upload["video_id"] = upload["video_id"].astype(str)
    up_full = upload["upload_date"].astype(str).str[:4]
    up = up_full[up_full.str.fullmatch(r"\d{4}")].astype(int)
    up.index = upload.loc[up.index, "video_id"]

    wb = openpyxl.load_workbook(f"{BASE}/data/ParkCeleb/PD_demo.xlsx",
                                read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip().lower().replace(" ", "") for c in rows[0]]
    yod = {}
    for r in rows[2:]:
        if r[0] is None:
            continue
        sp = str(r[0]).strip()
        yod[sp] = int(str(r[hdr.index("yod")]).strip())
    wb2 = openpyxl.load_workbook(f"{BASE}/data/ParkCeleb/speakers_pairs.xlsx",
                                 read_only=True)
    ws2 = wb2.active
    rows2 = list(ws2.iter_rows(values_only=True))
    pairs = {}
    for r in rows2[1:]:
        if r[0] is None:
            continue
        pairs[str(r[0]).strip()] = str(r[1]).strip()

    iv = np.full(len(X), np.nan)
    for i, sp in enumerate(sid):
        vid = kid[i]
        if vid not in up.index:
            continue
        uy = up[vid]
        if sp.startswith("pd_"):
            yr = yod.get(sp)
        else:
            yr = yod.get(pairs.get(sp))
        if yr is None:
            continue
        iv[i] = uy - yr

    out = {}
    for name, (lo, hi) in INTERVALS.items():
        sel = (iv >= lo) & (iv <= hi)
        out[name] = (X[sel], Y[sel], sid[sel], t_age_all[sel], t_sex_all[sel])
    return out

def gbm(**kw):
    return GradientBoostingClassifier(**kw)

def select_features(Xs, ys, seed):
    et = ExtraTreesClassifier(n_estimators=200, max_depth=10,
                              random_state=seed, n_jobs=-1).fit(Xs, ys)
    cols = np.argsort(et.feature_importances_)[::-1][:N_SELECT]
    return cols

def bin1(v, edges):
    return np.clip(np.searchsorted(edges, v, side="right") - 1,
                   0, len(edges) - 2)

def exact_pmf(names, vals, laplace=0.0):
    vals2 = vals.reshape(-1, len(names))
    K = vals2.shape[0]
    ncells = 1
    for n in names:
        ncells *= np.unique(vals2[:, names.index(n)]).size
    def _core(*args):
        arrs = [np.atleast_1d(a).reshape(-1) for a in args]
        Xm = np.column_stack(arrs)
        raw = np.array([np.mean((vals2 == row[None, :]).all(axis=1))
                        for row in Xm])
        return (raw * K + laplace) / (K + laplace * ncells)
    ns = {"_core": _core}
    exec(f"_pdf = lambda {','.join(names)}: _core({','.join(names)})", ns)
    return ns["_pdf"]

def source_cb_weights(Xsf, ys, ag, sx, up, gr):

    w = ca.causal_weights(
        causal_graph=UGAS_GRAPH, cause_var="y", effect_var="X",
        fit_method=CB_FIT_METHOD,
        X=Xsf,
        y=ys,
        Age=(ag, len(AGE_EDGES) - 1),
        Sex=sx,
        UPDRS=(up, len(UPDRS_EDGES) - 1),
        GRBAS=(gr, len(GRBAS_EDGES) - 1),
    )
    w = np.asarray(w)
    uy = np.unique(ys)
    col = {v: i for i, v in enumerate(uy)}
    return np.array([w[i, col[ys[i]]] for i in range(len(ys))])

def transport_ratio(ag, sx, t_age, t_sex):
    a_src = bin1(ag, AGE_EDGES).astype(float)
    a_tgt = bin1(t_age, AGE_EDGES).astype(float)
    s_src = sx.astype(float)
    s_tgt = np.round(np.asarray(t_sex, dtype=float))

    t_ok = ~(np.isnan(a_tgt) | np.isnan(s_tgt))
    p_tgt = exact_pmf(["Age", "Sex"],
                       np.column_stack([a_tgt[t_ok], s_tgt[t_ok]]),
                       laplace=0.5)
    p_src = exact_pmf(["Age", "Sex"],
                       np.column_stack([a_src, s_src]),
                       laplace=0.5)
    num = p_tgt(a_src, s_src)
    den = p_src(a_src, s_src)
    return num / den

def ugas_weights(Xsf, ys, ag, sx, up, gr, t_age, t_sex):
    w_source = source_cb_weights(Xsf, ys, ag, sx, up, gr)
    r = transport_ratio(ag, sx, t_age, t_sex)
    return w_source * r

def patient_metrics(y, proba, pids):
    del pids
    y = np.asarray(y).ravel()
    p = np.asarray(proba).ravel()
    ba = dur.balanced_acc(y, p)
    f1 = f1_score(y, p > 0.5, zero_division=0)
    auc = roc_auc_score(y, p) if len(np.unique(y)) > 1 else np.nan
    return ba, f1, auc

def grid_search(Xs, ys, pids, seed=0):
    from sklearn.model_selection import GroupKFold
    uniq = np.unique(pids)
    rng = np.random.RandomState(seed)
    best, best_s = None, -np.inf
    for ne in GBM_GRID["n_estimators"]:
        for md in GBM_GRID["max_depth"]:
            for lr in GBM_GRID["learning_rate"]:
                scores = []
                for rep in range(3):
                    order = rng.permutation(uniq)
                    folds = np.array_split(order, 5)
                    for fi in range(5):
                        tr_p = np.concatenate([f for j, f in enumerate(folds)
                                               if j != fi])
                        te_p = folds[fi]
                        tr = np.isin(pids, tr_p)
                        te = np.isin(pids, te_p)
                        m = GradientBoostingClassifier(
                            n_estimators=ne, max_depth=md, learning_rate=lr,
                            random_state=seed).fit(Xs[tr], ys[tr])
                        scores.append(roc_auc_score(ys[te], m.predict_proba(
                            Xs[te])[:, 1]))
                sc = np.mean(scores)
                if sc > best_s:
                    best_s, best = sc, {"n_estimators": ne, "max_depth": md,
                                        "learning_rate": lr}
    return best

def run_interval(name, Xs, ys, ag, sx, up, gr, pids, Xt, yt, sidt, t_age, t_sex):
    hp = grid_search(Xs, ys, pids, seed=RANDOM_SEED)
    print(f"    grid-search best params on source: {hp}")

    rng = np.random.RandomState(RANDOM_SEED)
    conf_ba, conf_f1, conf_auc = [], [], []
    dec_ba, dec_f1, dec_auc = [], [], []
    ess_runs = []
    for r in range(N_BOOT):
        seed = rng.randint(0, 2**31 - 1)
        cols = select_features(Xs, ys, seed)
        Xsf, Xtf = Xs[:, cols], Xt[:, cols]

        mod_conf = gbm(random_state=seed, **hp).fit(Xsf, ys)
        w = ugas_weights(Xsf, ys, ag, sx, up, gr, t_age, t_sex)
        ess_runs.append((w.sum() ** 2) / (w ** 2).sum())
        lo, hi = np.percentile(w, 5), np.percentile(w, 95)
        mod_deconf = gbm(random_state=seed + 1, **hp).fit(
            Xsf, ys, sample_weight=np.clip(w, lo, hi))

        p_conf = mod_conf.predict_proba(Xtf)[:, 1]
        p_deconf = mod_deconf.predict_proba(Xtf)[:, 1]
        cba, cf1, cauc = patient_metrics(yt, p_conf, sidt)
        dba, df1, dauc = patient_metrics(yt, p_deconf, sidt)
        conf_ba.append(cba); conf_f1.append(cf1); conf_auc.append(cauc)
        dec_ba.append(dba); dec_f1.append(df1); dec_auc.append(dauc)

    return {"BA conf": np.nanmedian(conf_ba), "BA deconf": np.nanmedian(dec_ba),
            "F1 conf": np.nanmedian(conf_f1), "F1 deconf": np.nanmedian(dec_f1),
            "AUC conf": np.nanmedian(conf_auc), "AUC deconf": np.nanmedian(dec_auc),
            "w_ess_med": float(np.median(ess_runs))}

def main():
    Xs, ys, ag, sx, up, gr, pids = load_nv_petaca()
    print(f"NV Petaca source: {Xs.shape[0]} recordings "
          f"({int(ys.sum())} PD / {len(ys) - int(ys.sum())} HC), "
          f"age [{ag.min():.0f}-{ag.max():.0f}], sex 1: {int(sx.sum())}, "
          f"UPDRS [{up.min():.0f}-{up.max():.0f}] (HC imputed 0), "
          f"GRBAS-G [0-{int(gr.max())}]")

    intervals = load_pc_intervals()
    for name, (X, y, s, ta, ts) in intervals.items():
        print(f"  interval {name:>2}: {X.shape[0]} clips, "
              f"{len(np.unique(s))} speakers, PD {int(y.sum())}, "
              f"age [{np.nanmin(ta):.0f}-{np.nanmax(ta):.0f}], "
              f"sex 1: {int(np.nansum(ts))}")

    print(f"\n{'int':<4}{'BA conf':>9}{'BA deconf':>10}{'F1 conf':>9}"
          f"{'F1 deconf':>10}{'AUC conf':>9}{'AUC deconf':>10}")
    results = {}
    for name, (X, y, s, ta, ts) in intervals.items():
        r = run_interval(name, Xs, ys, ag, sx, up, gr, pids, X, y, s, ta, ts)
        results[name] = r
        print(f"{name:<4}{r['BA conf']:>9.4f}{r['BA deconf']:>10.4f}"
              f"{r['F1 conf']:>9.4f}{r['F1 deconf']:>10.4f}"
              f"{r['AUC conf']:>9.4f}{r['AUC deconf']:>10.4f}")
        print(f"     dBA = {r['BA deconf'] - r['BA conf']:+.4f}   "
              f"w ESS (median of {N_BOOT}) = {r['w_ess_med']:.0f}")

if __name__ == "__main__":
    main()